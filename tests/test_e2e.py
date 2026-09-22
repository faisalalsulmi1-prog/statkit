"""Capstone end-to-end gate (PLAN Chunk 15).

For EVERY committed fixture, EVERY data sheet, and EVERY ``test_id`` in that
sheet's ``expect_block.target_tests`` (curation's intended tests), drive the FULL
pipeline the app drives:

    grid.load -> clean.clean -> infer.infer
      -> auto-bind the test's Contract roles from the inferred profiles
         (statkit.bind.auto_columns -- the SAME greedy contract-satisfying logic
          the Streamlit menu/pickers use)
      -> check.check  -> REGISTRY[test_id].run  (advise.advised-wrapped)
      -> sentences.render -> charts.chart -> report.build_report

DoD (PLAN §10 Chunk 15): "every corpus sheet reaches a result OR a stated block
for its intended test". So each (fixture, sheet, target_test) case must end in ONE
of two clean terminal states, never a crash:

  * BOUND & RAN -> a ``Result`` that is either populated (status "ok") or a
    ``status="blocked"`` Result carrying a stated, human-readable reason; and in
    both sub-cases: no exception, no NaN/inf statistic leaking as a live result, a
    non-empty rendered sentence, chart() that does not raise, and a Word report
    that round-trips (re-opens with python-docx).
  * GREYED OUT -> no column assignment satisfies the Contract for any accepted
    layout, so the app would grey the test out with a reason. That reason is the
    "stated block". We assert the reason is non-empty (a real, human-readable
    grey-out), and record the case.

Because ALL 39 data-sheet oracles carry ``served: true`` (none set
``block_reason``), the sheet-level DoD -- "the sheet reaches a result for its
intended test" -- is enforced by ``test_sheet_served_reaches_a_result``: at least
one of a served sheet's target tests must reach a populated non-blocked Result.
A sheet where NONE of its target tests can bind-and-run would be a real gap.

The NaN gate keys on the primary ``statistic``/``p``/``df`` (PLAN §5.2: gating is
on the test statistic/p, never on a NaN effect/OR -- McNemar b·c=0 yields a valid
p with a "not estimable" OR by design, snag 15).
"""
from __future__ import annotations

import functools
import io
import math
import re

import corpus
import numpy as np
import pandas as pd
import pytest
from docx import Document
from scipy.stats import rankdata

from statkit import bind, charts, check, clean, infer, report, sentences
from statkit.registry import REGISTRY

# case classifications, filled as cases run, harvested by the summary test.
_SEEN: dict[tuple[str, str, str], str] = {}


# --------------------------------------------------------------------------
# pipeline helpers (cached so 128 cases reuse 39 datasets and their results)
# --------------------------------------------------------------------------
def _grid_for(fid, sheet_name):
    grids = corpus.load_grids(fid)
    by_name = {g.sheet: g for g in grids}
    if sheet_name in by_name:
        return by_name[sheet_name]
    if len(grids) == 1:                 # csv/txt "(single)" vs the oracle's name
        return grids[0]
    return None


@functools.lru_cache(maxsize=None)
def _dataset(fid, sheet_name):
    grid = _grid_for(fid, sheet_name)
    assert grid is not None, f"no grid for {fid}/{sheet_name}"
    return infer.infer(clean.clean(grid))


def _params(spec) -> dict:
    """Registry defaults, exactly as the app seeds its controls; runners fill in a
    sensible success/reference/expected when these are None (props/cat/regress)."""
    return {p.name: p.default for p in spec.contract.params}


def _pick_layout(spec, dataset, oracle_layout):
    """Prefer the curation's intended layout when it binds, else the first that
    does (the app defaults to the first satisfiable layout)."""
    sat = bind.satisfiable_layouts(spec, dataset)
    if not sat:
        return None
    if oracle_layout in sat:
        return oracle_layout
    return sat[0]


@functools.lru_cache(maxsize=None)
def _run_case(fid, sheet_name, test_id, oracle_layout):
    """Drive the pipeline for one case, ending in one of three STATED terminal
    states (never a crash):

      * ("greyed", reason)      -- no binding satisfies the contract.
      * ("bind_error", reason)  -- the chosen test cannot bind this sheet's shape
        (``bind.BindError``, e.g. a count grid with fractional/negative cells).
        This mirrors the app's S14 handler, which catches BindError and shows the
        message as a stated error rather than a traceback (app.py step 7).
      * ("ran", result, spec, bound) -- bound and ran.

    Only ``bind.BindError`` is folded into a stated state; every OTHER exception
    (from ``run``/``sentences``/``charts``/``report``) still propagates so a
    genuine crash IS the RED signal in the owning parametrized test.
    """
    spec = REGISTRY[test_id]
    dataset = _dataset(fid, sheet_name)
    layout = _pick_layout(spec, dataset, oracle_layout)
    if layout is None:
        ok, why = bind.satisfies(spec, dataset)
        assert not ok
        return ("greyed", why)
    columns = bind.auto_columns(spec, dataset, layout)
    assert columns is not None, f"{fid}/{sheet_name}/{test_id}: satisfiable but no assignment"
    try:
        bound = bind.bind(spec, dataset, columns, layout=layout, params=_params(spec))
    except bind.BindError as exc:          # S14: a stated wrong-shape terminal state
        return ("bind_error", str(exc))
    profs = {p.name: p for p in dataset.profiles}
    bound.findings = check.check(bound, profs)
    result = spec.run(bound)
    return ("ran", result, spec, bound)


# --------------------------------------------------------------------------
# parametrization
# --------------------------------------------------------------------------
def _cases():
    out = []
    for fid in corpus.fixture_ids():
        for sh in corpus.expect(fid)["sheets"]:
            if sh.get("role") != "data":
                continue
            eb = sh["expect_block"]
            for tid in eb.get("target_tests", []):
                out.append((fid, sh["name"], tid, eb.get("layout")))
    return out


def _served_sheets():
    out = []
    for fid in corpus.fixture_ids():
        for sh in corpus.expect(fid)["sheets"]:
            if sh.get("role") != "data":
                continue
            eb = sh["expect_block"]
            if eb.get("served"):
                out.append((fid, sh["name"], tuple(eb.get("target_tests", [])),
                            eb.get("layout")))
    return out


_CASE_IDS = [f"{f}:{s}:{t}" for (f, s, t, _l) in _cases()]


# --------------------------------------------------------------------------
# assertions on a Result
# --------------------------------------------------------------------------
def _finite(x) -> bool:
    return isinstance(x, (int, float)) and math.isfinite(x)


def _assert_no_nan_statistic(result):
    """No NaN/inf leaking as a live statistic (PLAN §5.2 gating keys on stat/p)."""
    if result.statistic is not None:
        assert _finite(result.statistic[1]), \
            f"{result.test_id}: statistic {result.statistic} not finite in an ok result"
    if result.p is not None:
        assert _finite(result.p), f"{result.test_id}: p={result.p!r} not finite"
    for d in result.df:
        assert _finite(d), f"{result.test_id}: df {result.df} not finite"


def _assert_populated(result):
    """An ok Result must actually carry a computed answer, not an empty shell.
    (describe over categorical-only columns populates extra['categoricals']
    rather than a numeric descriptives frame -- still a real Table 1.)"""
    populated = (result.statistic is not None or result.p is not None
                 or result.estimate is not None or result.effect is not None
                 or result.descriptives is not None or result.table is not None
                 or bool(result.extra.get("categoricals")))
    assert populated, f"{result.test_id}: ok Result carries no statistic/table/estimate"


def _assert_stated_block(result):
    blocks = [f for f in result.findings if f.severity == "block"]
    assert blocks, f"{result.test_id}: blocked Result with no block finding"
    assert any(f.text.strip() for f in blocks), \
        f"{result.test_id}: block finding has empty reason text"


# --------------------------------------------------------------------------
# hardening assertions (this batch): pin the earlier waves' fixes end-to-end
# --------------------------------------------------------------------------
def _docx_text(doc) -> str:
    """All human-readable text of a python-docx Document: every paragraph
    (headings included) plus every table cell."""
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _block_reason(result) -> str:
    for f in result.findings:
        if f.severity == "block":
            return f.text
    return ""


def _human_readable(reason) -> bool:
    """A real, student-facing why: not empty, not a bare token. An internal
    placeholder ("measures__0"), an empty string, or a one-word code all fail."""
    r = (reason or "").strip()
    return len(r) >= 10 and " " in r and any(ch.islower() for ch in r)


def _real_names(dataset) -> set:
    """Every REAL name a role label may legitimately carry: a source column
    header, or a distinct level of a text/categorical column (a pivoted paired
    condition label such as 'Pre'/'Post'). The internal role placeholders
    'before'/'after'/'measures__i' are NOT in here unless the data truly uses
    them, which is exactly what pins B6."""
    names = {str(c) for c in dataset.df.columns}
    for c in dataset.df.columns:
        s = dataset.df[c]
        if s.dtype == object or str(s.dtype) == "string":
            names |= {str(v) for v in s.dropna().unique()}
    return names


# B6: internal, unambiguous placeholder that must NEVER reach a student surface.
_INTERNAL = "measures__"


def _assert_no_internal_names(result, sentence, captions, report_text, dataset):
    """Task 3 (pins B6): no internal role/derived-column name leaks to the
    student. 'measures__' is unambiguously internal -> banned on every surface.
    'before'/'after' are ordinary words, so we only forbid them AS the role
    LABEL when a real header/condition name should be there (checked against the
    dataset's real names, not by banning the words)."""
    surfaces = {"the sentence": sentence, "the report": report_text}
    for i, cap in enumerate(captions):
        surfaces[f"chart caption {i}"] = cap
    for g in result.groups:
        surfaces[f"groups[{g!r}]"] = str(g)
    for k, v in result.labels.items():
        surfaces[f"labels[{k!r}]"] = str(v)
    for where, text in surfaces.items():
        assert _INTERNAL not in text, (
            f"{result.test_id}: internal placeholder {_INTERNAL!r} leaked into "
            f"{where}: {text!r}")

    if result.test_id in ("t_paired", "wilcoxon"):
        real = _real_names(dataset)
        for role in ("before", "after"):
            lbl = result.labels.get(role)
            if lbl is not None:                    # one-sample wilcoxon has none
                assert str(lbl) in real, (
                    f"{result.test_id}: labels[{role!r}] = {lbl!r} is not a real "
                    "column header or condition level -- the internal role "
                    "placeholder leaked as a variable name (B6 regression)")


def _assert_mwu_direction(result):
    """Task 4 (pins B1): the stated 'higher' group is the one with the higher
    MEAN RANK, recomputed independently from the canonical frame -- an orthogonal
    check that would catch a regression to the old raw-means direction."""
    if result.test_id != "mwu":
        return
    frame = result.arrays.get("_frame")
    assert frame is not None and {"outcome", "group"} <= set(frame.columns), \
        f"mwu {result.test_id}: no canonical outcome+group frame to recompute from"
    order = list(dict.fromkeys(frame["group"].tolist()))
    assert len(order) == 2, f"mwu: expected 2 groups, got {order}"
    g0 = frame.loc[frame["group"] == order[0], "outcome"].to_numpy(float)
    g1 = frame.loc[frame["group"] == order[1], "outcome"].to_numpy(float)
    ranks = rankdata(np.concatenate([g0, g1]))      # pooled ranks, ties averaged
    mr0, mr1 = ranks[:len(g0)].mean(), ranks[len(g0):].mean()
    winner = str(order[0]) if mr0 > mr1 else str(order[1])   # runner's tie-break
    assert str(result.higher) == winner, (
        f"mwu direction: result.higher={result.higher!r} but the higher mean "
        f"rank is {winner!r} (mean ranks {mr0:.4f} vs {mr1:.4f}) -- raw-means bug?")


def _assert_cat_observed_named(result, sentence):
    """Task 5 (pins S4): chi2_ind / fisher carry the observed contingency table
    and name the actual row/column variables (never the generic phrase)."""
    if result.test_id not in ("chi2_ind", "fisher"):
        return
    obs = result.extra.get("observed")
    assert isinstance(obs, pd.DataFrame), (
        f"{result.test_id}: extra['observed'] contingency table is missing "
        f"(got {type(obs).__name__})")
    assert obs.shape[0] >= 2 and obs.shape[1] >= 2, \
        f"{result.test_id}: observed table is not at least 2x2: {obs.shape}"
    row, col = result.labels.get("row"), result.labels.get("col")
    assert row and col, \
        f"{result.test_id}: crosstab variables not named (S4): labels={result.labels}"
    assert "the two variables" not in sentence, \
        f"{result.test_id}: sentence used the generic 'the two variables' (S4 regression)"
    assert str(row) in sentence and str(col) in sentence, (
        f"{result.test_id}: sentence does not name both variables "
        f"({row!r}, {col!r}): {sentence!r}")


# Task 7 (pins B2): a p-value printed as EXACTLY zero. `fmt.p` collapses <.001 to
# "p < .001" and never emits these; a leak means fmt.p was bypassed. Precise so
# "p = 0.057" (a real robustness-check p) and "disp=0" (a function arg) do NOT
# trip it.
_P_ZERO = (
    re.compile(r"\bp\s*=\s*0(?![.\d])"),        # p = 0 / p=0 (bare zero)
    re.compile(r"\bp\s*=\s*0?\.0+(?!\d)"),      # p = .000 / 0.000 / 0.00 (all zeros)
)


def _assert_no_p_zero_or_nan(result, *texts):
    for text in texts:
        for pat in _P_ZERO:
            m = pat.search(text)
            assert m is None, (
                f"{result.test_id}: a p-value printed as exactly zero (B2 "
                f"regression): ...{text[max(0, m.start() - 15):m.start() + 15]!r}...")
        m = re.search(r"\bnan\b", text, re.I)   # \b so 'Banana'/'finance' are safe
        assert m is None, (
            f"{result.test_id}: 'nan'/'NaN' leaked into a user-facing surface: "
            f"...{text[max(0, m.start() - 15):m.start() + 15]!r}...")


# --------------------------------------------------------------------------
# the per-case gate
# --------------------------------------------------------------------------
@pytest.mark.parametrize("fid,sheet,tid,layout", _cases(), ids=_CASE_IDS)
def test_case_reaches_result_or_stated_block(fid, sheet, tid, layout):
    outcome = _run_case(fid, sheet, tid, layout)

    if outcome[0] == "greyed":
        _SEEN[(fid, sheet, tid)] = "greyed"
        assert outcome[1].strip(), f"{fid}/{sheet}/{tid}: greyed with no reason"
        return
    if outcome[0] == "bind_error":       # task 1: a stated wrong-shape terminal state
        _SEEN[(fid, sheet, tid)] = "bind_error"
        assert outcome[1].strip(), f"{fid}/{sheet}/{tid}: BindError with no message"
        return

    _, result, spec, bound = outcome
    dataset = _dataset(fid, sheet)
    assert result.status in ("ok", "blocked"), \
        f"{fid}/{sheet}/{tid}: unexpected status {result.status!r}"

    if result.status == "ok":
        _SEEN[(fid, sheet, tid)] = "ran_ok"
        _assert_populated(result)
        _assert_no_nan_statistic(result)
        _assert_mwu_direction(result)               # task 4 (pins B1)
    else:
        _SEEN[(fid, sheet, tid)] = "ran_blocked"
        _assert_stated_block(result)

    # a non-empty plain-English sentence in BOTH sub-cases
    sentence = sentences.render(result)
    assert isinstance(sentence, str) and sentence.strip(), \
        f"{fid}/{sheet}/{tid}: empty sentence"

    if result.status == "ok":
        _assert_cat_observed_named(result, sentence)   # task 5 (pins S4)

    # charts must not raise; each is (caption, PNG bytes). Blocked -> ().
    figs = charts.chart(result)
    assert isinstance(figs, tuple)
    for cap, png in figs:
        assert isinstance(cap, str)
        assert isinstance(png, (bytes, bytearray)) and png[:4] == b"\x89PNG", \
            f"{fid}/{sheet}/{tid}: chart {cap!r} is not PNG bytes"

    # the Word report must build AND round-trip (re-open with python-docx). The
    # dataset+bound are threaded so the S11 Data/provenance section is exercised
    # for every case (task 6), exactly as the app builds it.
    buf = report.build_report(result, spec, dataset=dataset, bound=bound)
    data = buf.getvalue()
    assert data[:2] == b"PK", f"{fid}/{sheet}/{tid}: report is not a .docx (zip)"
    report_doc = Document(io.BytesIO(data))          # raises if the docx is malformed
    report_text = _docx_text(report_doc)

    # task 3 (pins B6) + task 7 (pins B2): user-facing surfaces stay clean in BOTH
    # sub-cases (a blocked report's summary IS the block reason).
    _assert_no_internal_names(result, sentence, [c for c, _ in figs],
                              report_text, dataset)
    _assert_no_p_zero_or_nan(result, sentence, report_text)


# --------------------------------------------------------------------------
# the sheet-level DoD: a served sheet reaches a result for its intended test
# --------------------------------------------------------------------------
@pytest.mark.parametrize(
    "fid,sheet,targets,layout", _served_sheets(),
    ids=[f"{f}:{s}" for (f, s, _t, _l) in _served_sheets()])
def test_sheet_served_reaches_a_result(fid, sheet, targets, layout):
    """served:true means the sheet's intended test(s) reach a Result -- at least
    one target test must bind-and-run to a populated, non-blocked Result."""
    ran_ok = []
    for tid in targets:
        try:
            outcome = _run_case(fid, sheet, tid, layout)
        except Exception:               # a crash is caught by the per-case test
            continue
        if outcome[0] == "ran" and outcome[1].status == "ok":
            ran_ok.append(tid)
    assert ran_ok, (
        f"{fid}/{sheet}: served:true but NONE of {list(targets)} reached a "
        "populated result (all greyed/blocked/crashed) -- a real coverage gap")


# --------------------------------------------------------------------------
# coverage sanity + a printed summary (run with -s to see the breakdown)
# --------------------------------------------------------------------------
def test_e2e_coverage_summary():
    """Coverage tally + a guard (task 2): every non-ok case carries a real,
    human-readable reason, and the audited breakdown is PINNED so a case silently
    flipping ran_ok -> blocked/greyed/bind_error in the future trips this test."""
    from collections import Counter
    tally = Counter()
    greyed, blocked, bind_errors = [], [], []
    for fid, sheet, tid, layout in _cases():
        outcome = _run_case(fid, sheet, tid, layout)
        if outcome[0] == "greyed":
            tally["greyed"] += 1
            reason = outcome[1]
            greyed.append(f"{fid}/{sheet}/{tid}: {reason}")
            assert _human_readable(reason), \
                f"{fid}/{sheet}/{tid}: greyed reason not human-readable: {reason!r}"
        elif outcome[0] == "bind_error":
            tally["bind_error"] += 1
            reason = outcome[1]
            bind_errors.append(f"{fid}/{sheet}/{tid}: {reason}")
            assert _human_readable(reason), \
                f"{fid}/{sheet}/{tid}: bind_error reason not human-readable: {reason!r}"
        elif outcome[1].status == "ok":
            tally["ran_ok"] += 1
        else:
            tally["ran_blocked"] += 1
            reason = _block_reason(outcome[1])
            blocked.append(f"{fid}/{sheet}/{tid}: {reason}")
            assert _human_readable(reason), \
                f"{fid}/{sheet}/{tid}: block reason not human-readable: {reason!r}"
    total = sum(tally.values())
    print(f"\n[e2e] {total} cases: ran_ok={tally['ran_ok']} "
          f"ran_blocked={tally['ran_blocked']} greyed={tally['greyed']} "
          f"bind_error={tally['bind_error']}")
    for title, rows in (("BLOCKED (stated-reason)", blocked),
                        ("GREYED (cannot-bind)", greyed),
                        ("BIND_ERROR (wrong-shape, stated)", bind_errors)):
        if rows:
            print(f"[e2e] {title} cases:")
            for r in rows:
                print("   -", r)
    assert total == len(_cases())
    # Pin the audited 2026-09-21 breakdown. This is the tripwire for a silent
    # regression that turns a working case into a blocked/greyed/bind_error one
    # (or vice versa). Update these numbers DELIBERATELY when the corpus changes.
    assert tally["ran_ok"] == 111, f"ran_ok changed from 111: {tally['ran_ok']}"
    assert tally["ran_blocked"] == 14, f"ran_blocked changed from 14: {tally['ran_blocked']}"
    assert tally["greyed"] == 3, f"greyed changed from 3: {tally['greyed']}"
    assert tally["bind_error"] == 0, f"bind_error changed from 0: {tally['bind_error']}"


# --------------------------------------------------------------------------
# task 6 (pins S11): the report Data/provenance section round-trips with content
# --------------------------------------------------------------------------
def test_report_data_section_roundtrips():
    """With ``dataset``/``bound`` threaded, ``build_report`` emits a populated
    "Data" section (source, rows-read, header-row) that re-opens with python-docx.
    The per-case gate exercises the round-trip for every case; this pins the
    actual provenance CONTENT on a representative ran-ok case."""
    for fid, sheet, tid, layout in _cases():
        outcome = _run_case(fid, sheet, tid, layout)
        if outcome[0] != "ran" or outcome[1].status != "ok":
            continue
        dataset = _dataset(fid, sheet)
        if not (dataset.n_rows_read and dataset.header_row):
            continue                                   # need both facts to assert on
        _, result, spec, bound = outcome
        buf = report.build_report(result, spec, dataset=dataset, bound=bound)
        doc = Document(io.BytesIO(buf.getvalue()))     # round-trips or raises
        paras = [p.text for p in doc.paragraphs]
        assert "Data" in paras, \
            f"{fid}/{sheet}/{tid}: report has no Data heading (S11 not threaded)"
        assert f"Rows read (after cleaning): {dataset.n_rows_read}" in paras, \
            f"{fid}/{sheet}/{tid}: Data section missing the rows-read fact"
        assert f"Header row: {dataset.header_row}" in paras, \
            f"{fid}/{sheet}/{tid}: Data section missing the header-row fact"
        return
    pytest.fail("no ran-ok case carried header/rows provenance to check")
