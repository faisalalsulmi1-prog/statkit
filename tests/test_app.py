"""Chunk 14 — the Streamlit wizard, driven headless (PLAN §8 / §10 DoD).

Six scripted scenarios walk ``app.py`` end-to-end with ``AppTest`` on synthetic
sheets, asserting the app reaches a result + plain-English sentence + a Word
download (or a *stated block*) WITHOUT raising, that it holds all state in
``st.session_state``, and that it never calls ``st.cache_*``:

  a. Welch's independent-samples t-test on two groups
  b. one-way ANOVA on three groups (post-hoc table shown)
  c. chi-square test of independence on two categorical columns
  d. Pearson correlation on two numeric columns
  e. a BLOCKED case (constant outcome) -> the stated block reason, no crash
  f. a suggestion re-bind (shaky assumptions -> switch to the robust alternative)

The wizard is thin (orchestration only); every statistic lives in ``statkit``.
The menu greying and the column pickers both read the per-test Contract, so this
file also pins that a non-satisfiable test is disabled with a reason and that a
role picker offers only role-compatible columns.
"""
from __future__ import annotations

import io
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import statkit.report as report_mod

APP = Path(__file__).resolve().parent.parent / "app.py"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _xlsx(rows, sheet="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _app():
    at = AppTest.from_file(str(APP), default_timeout=90)
    at.run()
    return at


def _upload(at, data, name):
    at.file_uploader[0].set_value((name, data, MIME_XLSX))
    at.run()
    return at


def _set(at, key, value):
    at.get_by_key(key).set_value(value)
    at.run()
    return at


def _click(at, key):
    at.get_by_key(key).click()
    at.run()
    return at


def _texts(at):
    """All rendered text we might assert on (markdown / write / headers / alerts)."""
    out = []
    for coll in (at.markdown, at.header, at.subheader, at.title, at.caption,
                 at.info, at.warning, at.success, at.error):
        out.extend(m.value for m in coll)
    return out


def _has(at, needle):
    return any(needle.lower() in t.lower() for t in _texts(at))


# --------------------------------------------------------------------------
# scenario data
# --------------------------------------------------------------------------
def _two_group_sheet():
    rows = [["Score", "Group"]]
    rows += [[v, "A"] for v in (12, 14, 11, 13, 15, 10, 16, 12)]
    rows += [[v, "B"] for v in (20, 22, 19, 24, 21, 23, 18, 25)]
    return _xlsx(rows, "TwoGroups")


def _three_group_sheet():
    rows = [["Val", "Grp"]]
    rows += [[v, "G1"] for v in (10, 11, 9, 12, 10)]
    rows += [[v, "G2"] for v in (15, 16, 14, 15, 17)]
    rows += [[v, "G3"] for v in (20, 22, 19, 21, 20)]
    return _xlsx(rows, "ThreeGroups")


def _contingency_sheet():
    counts = {("Drug", "Cured"): 20, ("Drug", "NotCured"): 10,
              ("Placebo", "Cured"): 8, ("Placebo", "NotCured"): 22}
    rows = [["Treatment", "Outcome"]]
    for (tr, ou), n in counts.items():
        rows += [[tr, ou]] * n
    return _xlsx(rows, "Contingency")


def _correlation_sheet():
    # both columns continuous (>10 distinct, non-integer) so neither can serve as
    # a categorical group -> a 3+-group test genuinely cannot bind here.
    rows = [["X", "Y"]]
    for i in range(1, 16):
        x = i + 0.5
        y = 2.0 * x + (0.4 if i % 2 else -0.4)
        rows.append([x, round(y, 2)])
    return _xlsx(rows, "Corr")


def _constant_sheet():
    rows = [["Score", "Group"]]
    rows += [[5, "A"]] * 6 + [[5, "B"]] * 6
    return _xlsx(rows, "Constant")


def _skewed_two_group_sheet():
    rows = [["Score", "Group"]]
    rows += [[v, "A"] for v in (1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 50.0)]
    rows += [[v, "B"] for v in (2.0, 2.0, 2.0, 2.0, 2.0, 2.0, 60.0)]
    return _xlsx(rows, "Skewed")


# --------------------------------------------------------------------------
# (a) Welch's independent-samples t-test
# --------------------------------------------------------------------------
def test_scenario_a_welch_t():
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    _set(at, "col::outcome", "Score")
    _set(at, "col::group", "Group")
    _click(at, "run")
    assert at.exception == []
    # reaches a result + a plain-English sentence naming the exact variant
    assert _has(at, "Welch")
    assert _has(at, "t-test")
    # a downloadable Word report is offered
    assert len(at.download_button) >= 1
    # all wizard state lives in session_state
    assert len(at.session_state) > 0


# --------------------------------------------------------------------------
# (b) one-way ANOVA + post-hoc
# --------------------------------------------------------------------------
def test_scenario_b_oneway_anova_posthoc():
    at = _app()
    _upload(at, _three_group_sheet(), "threegroups.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::anova_1w")
    _set(at, "col::outcome", "Val")
    _set(at, "col::group", "Grp")
    _click(at, "run")
    assert at.exception == []
    assert _has(at, "ANOVA")
    # post-hoc table is shown (p < .05 in this fixture -> Tukey auto-runs)
    assert _has(at, "Tukey") or _has(at, "post-hoc") or _has(at, "Post-hoc")
    assert any(len(df.value) for df in at.dataframe) if at.dataframe else True
    assert len(at.download_button) >= 1


# --------------------------------------------------------------------------
# (c) chi-square test of independence
# --------------------------------------------------------------------------
def test_scenario_c_chi_square():
    at = _app()
    _upload(at, _contingency_sheet(), "contingency.xlsx")
    _set(at, "goal", "counts")
    _set(at, "pairing", "independent")
    _click(at, "pick::chi2_ind")
    _set(at, "col::row", "Treatment")
    _set(at, "col::col", "Outcome")
    _click(at, "run")
    assert at.exception == []
    assert _has(at, "Chi-square") or _has(at, "chi-square")
    assert len(at.download_button) >= 1


# --------------------------------------------------------------------------
# (d) Pearson correlation
# --------------------------------------------------------------------------
def test_scenario_d_pearson():
    at = _app()
    _upload(at, _correlation_sheet(), "corr.xlsx")
    _set(at, "goal", "relate")
    _click(at, "pick::pearson")
    _set(at, "col::x", "X")
    _set(at, "col::y", "Y")
    _click(at, "run")
    assert at.exception == []
    assert _has(at, "Pearson") or _has(at, "correlation")
    assert len(at.download_button) >= 1


# --------------------------------------------------------------------------
# (e) BLOCKED — constant outcome -> stated block reason, no crash
# --------------------------------------------------------------------------
def test_scenario_e_blocked_constant():
    at = _app()
    _upload(at, _constant_sheet(), "constant.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    _set(at, "col::outcome", "Score")
    _set(at, "col::group", "Group")
    at.run()
    assert at.exception == []
    # the block reason is surfaced (constant column) ...
    assert _has(at, "constant")
    # ... and Run is disabled, so no report is offered
    run_btn = at.get_by_key("run")
    assert run_btn.disabled is True
    assert len(at.download_button) == 0


# --------------------------------------------------------------------------
# (f) suggestion re-bind — shaky assumptions -> robust alternative
# --------------------------------------------------------------------------
def test_scenario_f_suggestion_rebind():
    at = _app()
    _upload(at, _skewed_two_group_sheet(), "skewed.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    _set(at, "col::outcome", "Score")
    _set(at, "col::group", "Group")
    _click(at, "run")
    assert at.exception == []
    # the parametric test ran, and a robust-alternative suggestion is offered
    assert _has(at, "t-test")
    assert _has(at, "Mann-Whitney") or _has(at, "robust") or _has(at, "rank")
    suggestion = at.get_by_key("suggest::0")
    assert suggestion is not None
    # one click switches to the robust test and re-runs on the same columns
    _click(at, "suggest::0")
    assert at.exception == []
    assert _has(at, "Mann-Whitney")
    assert len(at.download_button) >= 1


# --------------------------------------------------------------------------
# Contract-driven menu greying + role-compatible pickers (can't drift)
# --------------------------------------------------------------------------
def test_menu_greys_unsatisfiable_test_with_reason():
    at = _app()
    _upload(at, _correlation_sheet(), "corr.xlsx")   # two numeric columns only
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    # a 3+-group test cannot bind (no categorical group column) -> disabled
    anova_btn = at.get_by_key("pick::anova_1w")
    assert anova_btn.disabled is True
    assert isinstance(anova_btn.help, str) and anova_btn.help != ""


def test_pickers_offer_only_role_compatible_columns():
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    # the numeric outcome role offers Score, never the categorical Group
    outcome = at.get_by_key("col::outcome")
    assert "Score" in outcome.options
    assert "Group" not in outcome.options
    # the grouping role offers Group, never the numeric Score
    group = at.get_by_key("col::group")
    assert "Group" in group.options
    assert "Score" not in group.options


# --------------------------------------------------------------------------
# static L3 hygiene the DoD calls out explicitly
# --------------------------------------------------------------------------
def test_app_uses_session_state_and_no_cache():
    src = APP.read_text(encoding="utf-8")
    assert "session_state" in src
    assert "cache_data" not in src
    assert "cache_resource" not in src


# ==========================================================================
# BATCH: engine fixes wired into the UI (B3/B4/B5/S6/S7/S8/S9/S11/S14 + NICE).
# Each fixture/test below pins one wiring; the docstring names its expected RED.
# ==========================================================================
def _binary01_sheet():
    rows = [["Cured", "Group"]]
    rows += [[1, "A"], [0, "A"], [1, "A"], [0, "A"]]
    rows += [[1, "B"], [0, "B"], [1, "B"], [0, "B"]]
    return _xlsx(rows, "Binary01")


def _subtotal_sheet():
    rows = [["Name", "Score"],
            ["Alice", 10], ["Bob", 20],
            ["Total", 30],
            ["Carol", 15], ["Dave", 25]]
    return _xlsx(rows, "Subtot")


def _ambiguous_header_sheet():
    # two consecutive all-text rows -> the header detector is torn (ambiguous),
    # and each one names the columns differently, so a pick is observable.
    rows = [["Group", "Score"],
            ["Category", "Value"]]
    rows += [["A", 10], ["B", 12], ["A", 14], ["B", 9], ["A", 11], ["B", 13]]
    return _xlsx(rows, "Ambig")


def _ambiguous_thousands_sheet():
    rows = [["Amount"], ["1,234"], ["2,345"], ["3,456"]]
    return _xlsx(rows, "Amt")


def _spelling_sheet():
    rows = [["Sex", "Score"],
            ["Male", 1], ["male", 2], ["Male", 3],
            ["female", 4], ["Female", 5], ["female", 6]]
    return _xlsx(rows, "Spell")


def _mixed_counts_sheet():
    # two categorical columns AND a numeric count column -> chi2_ind can bind
    # BOTH the long layout (row/col) and the table layout (counts) -> 2 layouts.
    rows = [["Treatment", "Outcome", "Count"],
            ["Drug", "Cured", 20], ["Drug", "NotCured", 10],
            ["Placebo", "Cured", 8], ["Placebo", "NotCured", 22]]
    return _xlsx(rows, "Mixed")


def _one_col_numeric_sheet():
    rows = [["Weight"]] + [[v] for v in (70, 72, 68, 75, 71, 69, 73, 74)]
    return _xlsx(rows, "W")


# --- B3: level dropdowns show display strings ("1", never "1.0") -----------
def test_b3_level_choice_shows_int_not_float():
    """Expected RED: the success-level dropdown for a 0/1 column offers
    '1.0'/'0.0' (the raw float str), so '1' is absent."""
    at = _app()
    _upload(at, _binary01_sheet(), "bin.xlsx")
    _set(at, "goal", "counts")
    _click(at, "pick::prop_1")
    _set(at, "col::outcome", "Cured")
    opts = at.get_by_key("param::success").options
    assert "1" in opts and "0" in opts
    assert "1.0" not in opts and "0.0" not in opts


# --- B4: restore wrongly-dropped rows -------------------------------------
def test_b4_restore_dropped_row_grows_data():
    """Expected RED: there is no 'restore_rows' control, so setting it raises
    KeyError (the widget does not exist yet)."""
    at = _app()
    _upload(at, _subtotal_sheet(), "subtot.xlsx")
    n_before = len(at.dataframe[0].value)
    _set(at, "restore_rows", [4])          # bring the dropped "Total" row back
    assert at.exception == []
    assert len(at.dataframe[0].value) == n_before + 1


# --- B5: header picker is stable (sticks across reruns) -------------------
def test_b5_header_pick_sticks_across_reruns():
    """Pins that a chosen header row persists and is reflected in the columns
    across a couple of run cycles (no oscillation)."""
    at = _app()
    _upload(at, _ambiguous_header_sheet(), "ambig.xlsx")
    _set(at, "header_row", 2)               # pick the 2nd candidate row
    assert at.get_by_key("header_row").value == 2
    assert "Category" in list(at.dataframe[0].value.columns)
    at.run()                               # a bare rerun must not revert it
    assert at.get_by_key("header_row").value == 2
    assert "Category" in list(at.dataframe[0].value.columns)


# --- S6: number-convention override ---------------------------------------
def test_s6_convention_override_changes_reading():
    """Expected RED: there is no per-column 'conv::' control, so setting it
    raises KeyError (the widget does not exist yet)."""
    at = _app()
    _upload(at, _ambiguous_thousands_sheet(), "amt.xlsx")
    auto = [float(x) for x in at.dataframe[0].value["Amount"].dropna().tolist()]
    assert max(auto) > 1000                 # auto: "1,234" -> 1234 (thousands)
    sb = at.get_by_key("conv::Amount")
    eu = [o for o in sb.options if o.startswith("1.234")][0]
    _set(at, "conv::Amount", eu)            # force dot-thousands, comma-decimal
    assert at.exception == []
    forced = [float(x) for x in at.dataframe[0].value["Amount"].dropna().tolist()]
    assert max(forced) < 10                 # forced: "1,234" -> 1.234


# --- S7: level-merge control ----------------------------------------------
def test_s7_level_merge_collapses_spellings():
    """Expected RED: there is no 'merge::' control, so setting it raises
    KeyError (the widget does not exist yet)."""
    at = _app()
    _upload(at, _spelling_sheet(), "spell.xlsx")
    assert at.dataframe[0].value["Sex"].nunique() == 4
    _set(at, "merge::Sex", True)
    assert at.exception == []
    assert at.dataframe[0].value["Sex"].nunique() == 2


# --- S8: search the test menu by name / alias -----------------------------
def _menu_keys(at):
    return {b.key for b in at.button if b.key and b.key.startswith("pick::")}


def test_s8_search_surfaces_welch():
    """Expected RED: there is no 'test_search' box, so setting it raises
    KeyError (the widget does not exist yet)."""
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "test_search", "welch")        # goal defaults to 'describe'
    assert "pick::t_ind" in _menu_keys(at)


def test_s8_search_surfaces_shapiro():
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "test_search", "shapiro")
    assert "pick::normality" in _menu_keys(at)


# --- S9: friendly layout labels -------------------------------------------
def test_s9_layout_labels_are_friendly():
    """Expected RED: the layout selectbox shows raw ids ('long'/'table')."""
    at = _app()
    _upload(at, _mixed_counts_sheet(), "mixed.xlsx")
    _set(at, "goal", "counts")
    _set(at, "pairing", "independent")
    _click(at, "pick::chi2_ind")
    sb = at.get_by_key("layout")
    assert len(sb.options) > 1
    assert "long" not in sb.options and "table" not in sb.options
    assert any("row" in o.lower() or "count" in o.lower() for o in sb.options)


# --- S11: the Word report carries the Data provenance section --------------
def test_s11_report_threads_dataset_and_bound():
    """Expected RED: _render_result calls build_report(result, spec) with no
    dataset/bound, so the report has no Data section."""
    captured = {}
    orig = report_mod.build_report

    def spy(result, spec=None, dataset=None, bound=None):
        buf = orig(result, spec, dataset=dataset, bound=bound)
        captured["dataset"] = dataset
        captured["bound"] = bound
        captured["bytes"] = buf.getvalue()
        return io.BytesIO(captured["bytes"])

    report_mod.build_report = spy
    try:
        at = _app()
        _upload(at, _two_group_sheet(), "twogroups.xlsx")
        _set(at, "goal", "compare")
        _set(at, "pairing", "independent")
        _click(at, "pick::t_ind")
        _set(at, "col::outcome", "Score")
        _set(at, "col::group", "Group")
        _click(at, "run")
        assert at.exception == []
    finally:
        report_mod.build_report = orig

    assert captured.get("dataset") is not None
    assert captured.get("bound") is not None
    doc = Document(io.BytesIO(captured["bytes"]))
    headings = [p.text for p in doc.paragraphs
                if p.style is not None and p.style.name.startswith("Heading")]
    assert "Data" in headings
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Rows read" in text or "Header row" in text


# --- S14: friendly errors, no traceback -----------------------------------
def test_s14_friendly_error_on_garbage_upload():
    """Expected RED: grid.load raises GridError uncaught -> at.exception set."""
    at = _app()
    _upload(at, b"this is definitely not a spreadsheet", "junk.xlsx")
    assert at.exception == []
    assert _has(at, "Save As") or _has(at, "readable") or _has(at, "couldn't")


# --- NICE: per-upload state reset -----------------------------------------
def test_nice_new_upload_clears_stale_test():
    """Expected RED: uploading a new sheet leaves the old test_id selected
    (goal/pairing persist, so the menu logic alone does not clear it)."""
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    assert at.session_state.get("test_id") == "t_ind"
    _upload(at, _three_group_sheet(), "threegroups.xlsx")
    assert at.session_state.get("test_id") is None


# --- NICE: stale-result reset on a changed selection ----------------------
def test_nice_changed_param_hides_previous_result():
    """Expected RED: after a Run, toggling a param keeps _ran True, so the old
    result (and its download button) stays on screen."""
    at = _app()
    _upload(at, _two_group_sheet(), "twogroups.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_ind")
    _set(at, "col::outcome", "Score")
    _set(at, "col::group", "Group")
    _click(at, "run")
    assert len(at.download_button) >= 1
    _set(at, "param::equal_var", True)      # change the selection
    assert len(at.download_button) == 0     # result hidden until re-run


# --- NICE: mu0 required for a one-sample t-test ----------------------------
def test_nice_mu0_required_for_one_sample():
    """Expected RED: mu0 silently defaults to 0.0, so Run is enabled with no
    reference value entered."""
    at = _app()
    _upload(at, _one_col_numeric_sheet(), "w.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")
    _click(at, "pick::t_1s")
    _set(at, "col::outcome", "Weight")
    assert at.get_by_key("run").disabled is True    # no mu0 yet -> disabled
    assert len(at.download_button) == 0
    _set(at, "param::mu0", 70.0)
    assert at.get_by_key("run").disabled is False


# ==========================================================================
# BATCH 2: five confirmed defect fixes (F3 / F8 / F9 / F11 / F12).
# ==========================================================================
def _skewed_one_col_sheet():
    rows = [["Weight"]] + [[v] for v in
            (1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 50.0)]
    return _xlsx(rows, "Skew1")


def _before_after_sheet():
    rows = [["Before", "After"]]
    rows += [[b, a] for b, a in
             ((10, 12), (8, 11), (9, 14), (7, 10), (11, 13), (6, 9))]
    return _xlsx(rows, "BA")


def _paired_long_dup_sheet():
    rows = [["Subject", "Condition", "Score"],
            ["S1", "Before", 10], ["S1", "After", 12], ["S1", "After", 13],
            ["S2", "Before", 8], ["S2", "After", 11],
            ["S3", "Before", 9], ["S3", "After", 14],
            ["S4", "Before", 7], ["S4", "After", 10]]
    return _xlsx(rows, "PairedDup")


def _regression_mixed_sheet():
    ages = (21.4, 34.1, 45.7, 29.3, 52.8, 38.2, 41.6, 33.9,
            60.5, 27.1, 48.4, 55.2, 31.7, 44.3)
    cities = ("NY", "LA", "SF")
    rows = [["Outcome", "Age", "City"]]
    for i, a in enumerate(ages):
        rows.append([round(2.0 * a + (3 if i % 2 else -3), 1), a, cities[i % 3]])
    return _xlsx(rows, "Reg")


# --- F3: a one-sample suggestion must carry its pairing (not dead-end) ------
def test_f3_one_sample_suggestion_switches_pairing_and_carries_mu0():
    """Expected RED: _apply_suggestion sets test_id but not pairing, so the
    Wilcoxon suggestion (paired) is filtered out when the student answered
    'independent' for the one-sample test -> 'Choose a test above.'"""
    at = _app()
    _upload(at, _skewed_one_col_sheet(), "skew1.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")     # the natural answer for a 1-sample test
    _click(at, "pick::t_1s")
    _set(at, "col::outcome", "Weight")
    _set(at, "param::mu0", 5.0)            # required for t_1s
    _click(at, "run")
    assert at.exception == []
    assert _has(at, "Wilcoxon") or _has(at, "rank")   # the robust suggestion shows
    assert at.get_by_key("suggest::0") is not None
    _click(at, "suggest::0")
    assert at.exception == []
    assert _has(at, "Wilcoxon")                       # a result renders ...
    assert not _has(at, "Choose a test above")        # ... it did NOT dead-end
    assert at.session_state["param::mu0"] == 5.0      # and mu0 carried across


# --- F8: one-sample Wilcoxon must require mu0 (no silent test against 0) ----
def test_f8_one_sample_wilcoxon_requires_mu0():
    """Expected RED: mu0 is not required for Wilcoxon, so the one-sample mode
    silently tests against mu0 = 0 with Run enabled."""
    at = _app()
    _upload(at, _one_col_numeric_sheet(), "w.xlsx")
    _set(at, "test_search", "wilcoxon")
    _click(at, "pick::wilcoxon")
    _set(at, "col::outcome", "Weight")
    assert at.get_by_key("run").disabled is True      # no mu0 -> held, not silent 0
    assert len(at.download_button) == 0
    _set(at, "param::mu0", 65.0)
    assert at.get_by_key("run").disabled is False


def test_f8_paired_wilcoxon_needs_no_mu0():
    """Guard: the paired (before/after) Wilcoxon needs no mu0 -> Run is enabled."""
    at = _app()
    _upload(at, _before_after_sheet(), "ba.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "paired")
    _click(at, "pick::wilcoxon")
    _set(at, "layout", "wide")
    _set(at, "col::before", "Before")
    _set(at, "col::after", "After")
    assert at.get_by_key("run").disabled is False


# --- F9: the duplicate-measurement warning must reach the student ----------
def test_f9_duplicate_measurement_warning_reaches_report():
    """Expected RED: app.py overwrites bind's DUP finding with check.check, so
    the 'duplicate measurement(s) ignored' warning never reaches the student."""
    captured = {}
    orig = report_mod.build_report

    def spy(result, spec=None, dataset=None, bound=None):
        buf = orig(result, spec, dataset=dataset, bound=bound)
        captured["bytes"] = buf.getvalue()
        captured["findings"] = tuple(result.findings)
        return io.BytesIO(captured["bytes"])

    report_mod.build_report = spy
    try:
        at = _app()
        _upload(at, _paired_long_dup_sheet(), "pd.xlsx")
        _set(at, "goal", "compare")
        _set(at, "pairing", "paired")
        _click(at, "pick::t_paired")
        _set(at, "col::subject", "Subject")
        _set(at, "col::condition", "Condition")
        _set(at, "col::outcome", "Score")
        _click(at, "run")
        assert at.exception == []
    finally:
        report_mod.build_report = orig

    assert any(f.code == "DUP" for f in captured["findings"])
    text = "\n".join(p.text for p in Document(io.BytesIO(captured["bytes"])).paragraphs)
    assert "duplicate measurement" in text.lower()


def test_f9_non_duplicate_sheet_findings_unchanged():
    """Guard: a clean paired sheet (no dup) carries NO DUP finding -> the merge
    reduces to check.check's findings exactly (composition unchanged)."""
    captured = {}
    orig = report_mod.build_report

    def spy(result, spec=None, dataset=None, bound=None):
        captured["findings"] = tuple(result.findings)
        return orig(result, spec, dataset=dataset, bound=bound)

    report_mod.build_report = spy
    try:
        at = _app()
        rows = [["Subject", "Condition", "Score"],
                ["S1", "Before", 10], ["S1", "After", 12],
                ["S2", "Before", 8], ["S2", "After", 11],
                ["S3", "Before", 9], ["S3", "After", 14],
                ["S4", "Before", 7], ["S4", "After", 10]]
        _upload(at, _xlsx(rows, "Clean"), "clean.xlsx")
        _set(at, "goal", "compare")
        _set(at, "pairing", "paired")
        _click(at, "pick::t_paired")
        _set(at, "col::subject", "Subject")
        _set(at, "col::condition", "Condition")
        _set(at, "col::outcome", "Score")
        _click(at, "run")
        assert at.exception == []
    finally:
        report_mod.build_report = orig
    assert not any(f.code == "DUP" for f in captured["findings"])


# --- F11: raw Python exception text must not reach the student -------------
def test_f11_friendly_error_when_run_raises():
    """Expected RED: the generic handler shows f'...: {exc}', leaking the raw
    exception text to the student."""
    from statkit.registry import REGISTRY
    orig = REGISTRY["t_ind"].run

    def boom(bound):
        raise ValueError("keywords must be strings")

    object.__setattr__(REGISTRY["t_ind"], "run", boom)
    try:
        at = _app()
        _upload(at, _two_group_sheet(), "twogroups.xlsx")
        _set(at, "goal", "compare")
        _set(at, "pairing", "independent")
        _click(at, "pick::t_ind")
        _set(at, "col::outcome", "Score")
        _set(at, "col::group", "Group")
        _click(at, "run")
    finally:
        object.__setattr__(REGISTRY["t_ind"], "run", orig)
    assert at.exception == []
    assert not _has(at, "keywords must be strings")
    assert _has(at, "could not be computed") or _has(at, "different test")


def test_f11_friendly_error_when_bind_raises():
    """Expected RED: the prepare-stage handler shows f'...: {exc}', leaking the
    raw exception text to the student."""
    import statkit.bind as bindmod
    orig = bindmod.bind

    def boom(*a, **k):
        raise RuntimeError("keywords must be strings")

    bindmod.bind = boom
    try:
        at = _app()
        _upload(at, _two_group_sheet(), "twogroups.xlsx")
        _set(at, "goal", "compare")
        _set(at, "pairing", "independent")
        _click(at, "pick::t_ind")
        _set(at, "col::outcome", "Score")
        _set(at, "col::group", "Group")
        at.run()
    finally:
        bindmod.bind = orig
    assert at.exception == []
    assert not _has(at, "keywords must be strings")
    assert _has(at, "prepare") or _has(at, "different test")


# --- F12: reference-level control skips a numeric predictor ----------------
def test_f12_reference_dropdown_skips_numeric_predictor():
    """Expected RED: the reference control uses predictors[0] regardless of
    type, so a numeric first predictor is offered its numbers as levels; there
    is no per-categorical-predictor 'param::reference::<header>' control."""
    at = _app()
    _upload(at, _regression_mixed_sheet(), "reg.xlsx")
    _set(at, "goal", "predict")
    _click(at, "pick::ols_multi")
    _set(at, "col::outcome", "Outcome")
    _set(at, "cols::predictors", ["Age", "City"])     # numeric predictor FIRST
    sb = at.get_by_key("param::reference::City")
    assert sb is not None
    assert "NY" in sb.options
    assert all(str(a) not in sb.options for a in (21.4, 34.1, 45.7))


# ==========================================================================
# BATCH 3 (W7): four confirmed app fixes (S-A / S-C / S-E + a dead import).
# ==========================================================================
def _skewed_two_numeric_sheet():
    # Weight is skewed (drives the one-sample t-test's normality flag -> the robust
    # one-sample Wilcoxon suggestion); Other is a SECOND numeric column so the
    # suggested Wilcoxon also has a satisfiable wide (before/after) layout, making
    # the layout selectbox appear and default to "Wide" — the dead-end this fixes.
    skew = (1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 50.0)
    other = (3.0, 6.0, 4.0, 8.0, 5.0, 9.0, 2.0, 7.0, 10.0, 12.0)
    rows = [["Weight", "Other"]] + [[w, o] for w, o in zip(skew, other)]
    return _xlsx(rows, "Skew2")


def _categorical_only_sheet():
    rows = [["Sex", "City"],
            ["Male", "NY"], ["Female", "LA"], ["Male", "SF"],
            ["Female", "NY"], ["Male", "LA"], ["Female", "SF"]]
    return _xlsx(rows, "Cats")


def _linear_two_col_sheet():
    rows = [["Y", "X"]]
    for i in range(1, 15):
        x = float(i)
        rows.append([round(2.0 * x + (0.5 if i % 2 else -0.5), 2), x])
    return _xlsx(rows, "Lin")


# --- S-A: a one-sample-Wilcoxon suggestion on a two-numeric sheet must not
#          dead-end on the Wide (before/after) layout --------------------------
def test_sa_wilcoxon_suggestion_lands_on_one_sample_not_wide():
    """Expected RED: _apply_suggestion carries test_id + pairing but NOT layout,
    so on a sheet with two numeric columns the suggested Wilcoxon's layout box
    defaults to 'Wide' -> empty before/after pickers -> no result, no download."""
    at = _app()
    _upload(at, _skewed_two_numeric_sheet(), "skew2.xlsx")
    _set(at, "goal", "compare")
    _set(at, "pairing", "independent")   # the natural answer for a one-sample test
    _click(at, "pick::t_1s")
    _set(at, "col::outcome", "Weight")
    _set(at, "param::mu0", 5.0)          # required for t_1s
    _click(at, "run")
    assert at.exception == []
    assert _has(at, "Wilcoxon") or _has(at, "rank")   # the robust suggestion shows
    assert at.get_by_key("suggest::0") is not None
    _click(at, "suggest::0")
    assert at.exception == []
    # lands on a POPULATED one-sample Wilcoxon result, not the before/after dead-end
    assert at.session_state.get("layout") == "one_sample"
    assert _has(at, "Wilcoxon")
    assert not _has(at, "Choose a test above")
    assert len(at.download_button) >= 1               # a real result + report
    assert at.session_state["param::mu0"] == 5.0      # μ₀ carried across


# --- S-C: describe on categorical columns must render the count tables --------
def test_sc_describe_categoricals_render_when_no_numeric():
    """Expected RED: _render_result never shows result.extra['categoricals'], so a
    describe on categorical-only columns (descriptives is None) shows an empty
    panel — the count tables never reach the student."""
    at = _app()
    _upload(at, _categorical_only_sheet(), "cats.xlsx")
    _set(at, "goal", "describe")
    _click(at, "pick::describe")
    _set(at, "cols::variables", ["Sex", "City"])
    _click(at, "run")
    assert at.exception == []
    # the categorical count tables each carry a 'Count' column; the "what I read"
    # preview df has columns Sex/City but no 'Count' column, so this isolates them.
    count_dfs = [d.value for d in at.dataframe
                 if "Count" in [str(c) for c in d.value.columns]]
    assert count_dfs, "no categorical count table rendered in the result panel"
    seen = set()
    for df in count_dfs:
        seen.update(str(v) for v in df.astype(str).values.ravel())
    assert {"Male", "Female"} <= seen
    assert "NY" in seen


# --- S-E: perfect-fit regression -> the panel F must match report/sentence ----
def test_se_perfect_fit_panel_renders_dash_not_astronomical_f():
    """Expected RED: _render_numbers prints fmt.num(F); on a perfect fit F is a
    huge FINITE number, so the panel shows a 30-digit F while the report and
    sentence (W6) show '—'. The panel must render '—' too."""
    import math as _math
    from statkit.registry import REGISTRY
    orig = REGISTRY["ols_simple"].run

    def perfect(bound):
        r = orig(bound)                  # a real, renderable OK result ...
        r.statistic = ("F", 1e30)        # ... forced to a perfect-fit shape:
        r.effect = ("f2", _math.inf)     # non-finite f² == perfect fit
        return r

    object.__setattr__(REGISTRY["ols_simple"], "run", perfect)
    try:
        at = _app()
        _upload(at, _linear_two_col_sheet(), "lin.xlsx")
        _set(at, "goal", "predict")
        _click(at, "pick::ols_simple")
        _set(at, "col::outcome", "Y")
        _set(at, "col::x", "X")
        _click(at, "run")
    finally:
        object.__setattr__(REGISTRY["ols_simple"], "run", orig)
    assert at.exception == []
    panel = "\n".join(m.value for m in at.markdown)
    assert "1000000000" not in panel     # no astronomical F number in the panel
    assert "**F** = —" in panel          # rendered as the em dash, like the report


# --- the unused BytesIO import is gone -----------------------------------------
def test_app_has_no_unused_bytesio_import():
    """Expected RED: app.py imports BytesIO but never uses it (Fable-confirmed)."""
    src = APP.read_text(encoding="utf-8")
    assert "BytesIO" not in src


# --- credit/contact footer shows on every view ---------------------------------
def test_footer_credit_shows_on_bare_run():
    """The credit/contact footer renders on EVERY view — even before any upload,
    where main() returns early — so it must sit at the page's very bottom."""
    at = _app()
    blocks = [m.value for m in at.caption] + [m.value for m in at.markdown]
    assert any("Made by Faisal Alsulami" in b and "faisal.alsulmi1@gmail.com" in b
               for b in blocks), \
        f"footer credit (name + contact email) missing on bare run; blocks={blocks!r}"
