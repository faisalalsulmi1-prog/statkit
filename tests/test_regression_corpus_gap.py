"""Regression pins for the "corpus-gap family" fixes (Fix Waves #1-#2).

These three sheets are the *next members* of the corpus-gap family that the
committed fixture corpus (tests/fixtures/generated/, expect=35) happens not to
contain. Each is built IN-MEMORY, driven through the SAME pipeline the app drives
(grid -> clean -> infer -> bind -> check -> run -> sentence -> chart -> report),
and asserts the FIXED student-facing behaviour so a silent regression trips here.

RED-first note (the fixes are already shipped, so these are green now). For each
pin, the pre-fix output is stated in the assertion's own comment; every pin was
verified to be genuinely exercised (mechanism + a simulated pre-fix string that
the matcher catches), not vacuously green:

  (a) all-text survey  -> NB1 (units row eats first data row), NB2 (Likert TEXT
      crashes a rank test), S-H (all-text sheet flagged header-ambiguous).
  (b) numeric-coded long-paired -> S-B(a)/F4: condition labels render "1"/"2",
      never the float "1.0"/"2.0", after the long->wide pivot.
  (c) die-face goodness-of-fit  -> S-B(b)/F4: face labels render "1".."6",
      never "1.0".."6.0".

Wave #3 (the corpus-gap SEAM defects, this file's new pins). Each line: test ->
the exact WRONG observable current (unfixed) code produces -> the assertion that
catches it. These 15 tests FAIL against current code (RED-first); the first
failing assertion of each is named so "my matcher is usually the bug" cannot bite.

  T1 trailing-blank header cell : clean picks the first DATA row as the header
     (header_row==2, header==['M','Yes','North','Placebo'], 22 rows, excel_rows[0]
     ==3; chi2_ind then binds row='M'/col='Yes'). CATCH: `t.header_row == 1`
     (assert 2 == 1) -> then header names / 24 rows / chi2 labels Sex/Smoker / n=24.
  T2 leading-blank header cell  : same eaten-header defect (header_row==2, 22
     rows). CATCH: `t.header_row == 1` (assert 2 == 1) -> ['Column A',...], 24 rows.
  T3 sub-1 doses merge          : the CATEGORICAL df column is built via fmt.num
     (2dp) so 0.001 & 0.002 both render '0' and MERGE, though the profile keeps 4
     levels. groups==('0','0.01','control'), n['0']==12. CATCH: `result.groups ==
     ('0.001','0.002','0.005','control')` (merged-groups tuple mismatch).
  T4 rounded-cell t-test        : a mixed numeric+junk column is stored 2dp then
     re-parsed, so t is computed on ROUNDED values: t==-8.878 (true -13.6626...).
     CATCH: `abs(result.statistic[1] - ss.ttest_ind(a,b,equal_var=False).stat)<1e-9`.
  T5x3 binary-text correlate    : Passed (Yes/No) is never coded, so x/y.to_numpy
     (float) raises ValueError: could not convert string to float: 'No' inside
     spec.run. CATCH: the propagated ValueError (crash pin) -> then r==0.81267 /
     tau==0.69206 / slope==18.0, coded set=={0.,1.}, BINCODE finding names it.
  T6 satisfaction vocab         : the 5-point satisfaction scale is not a known
     ORDERED_VOCAB, so kinds==(CATEGORICAL,) (and mwu cannot bind an outcome).
     CATCH: `prof.kinds[0] == Kind.ORDINAL` (assert CATEGORICAL == ORDINAL).
  T7 override on known vocab     : override->ORDINAL sends low/medium/high down the
     numeric-family path -> level_order None, levels (), false '20 text cell(s)
     treated as missing' note, then ValueError: 'low' inside spec.run. CATCH: the
     propagated ValueError (crash pin) -> then level_order/n_levels/no-missing-note.
  T8 override on unknown scale   : same numeric-family path on Terrible/Bad/OK/Great
     -> false missing note + ValueError: 'Terrible' inside spec.run. CATCH: the
     propagated ValueError (crash pin) -> then 4 levels + ORDORDER block + blocked.
  T9 fractional near-contiguous  : a 22-distinct 1-dp BMI column is mis-typed ID
     (near-contiguity ignores the fractional part), so t_ind sees "no numeric
     column". kinds==(ID,). CATCH: `prof.kinds == (Kind.NUMERIC, Kind.ID)`.
  T10 honest no-numeric message  : two integer near-contiguous columns are both ID
     and satisfies(t_paired) returns the LIE 'needs a numeric column; your sheet
     has none'. CATCH: `'Before' in why and 'After' in why` (+ no 'has none').
  T11 Fisher R×C scalar cast     : the Monte-Carlo Fisher pvalue is an array;
     float(res.pvalue) raises TypeError: only 0-dimensional arrays... at cat.py.
     CATCH: the propagated TypeError (crash pin) -> then ok / 0<=p<=1 / docx.
  T12 NA-token data row eaten     : a real data row of NA tokens (Ali,NA,n/a,na)
     passes the units-row shape gate and is CONSUMED as units (units_row==2, 5
     rows, silent). CATCH: `t.units_row is None` (assert 2 is None) -> 6 rows.
  T13 partial auto-bind           : auto_columns(spearman) returns the partial
     {'x':('Score',),'y':()} instead of None when a required role has no column.
     CATCH: `cols is None`.

Purely additive: no source, no existing test, and no committed oracle is touched.
"""
from __future__ import annotations

import io
import re

import pandas as pd
import pytest
import scipy.stats as ss
from docx import Document

from statkit import bind, charts, check, clean, grid, infer, report, sentences
from statkit.model import Kind
from statkit.registry import REGISTRY


# --------------------------------------------------------------------------
# in-memory pipeline driver (no disk, no fixture files -- ponytail: CSV bytes)
# --------------------------------------------------------------------------
def _csv(rows: list[str]) -> bytes:
    return ("\n".join(rows)).encode()


class _Run:
    """Every stage a student's answer passes through, kept for assertions."""

    def __init__(self, csv_bytes: bytes, name: str, test_id: str, layout: str,
                 overrides=None):
        grids = grid.load(csv_bytes, name)
        assert len(grids) == 1, f"{name}: expected one sheet, got {len(grids)}"
        self.table = clean.clean(grids[0])
        self.dataset = infer.infer(self.table, overrides=overrides)
        self.spec = REGISTRY[test_id]
        self.columns = bind.auto_columns(self.spec, self.dataset, layout)
        assert self.columns is not None, f"{name}: {test_id} could not auto-bind"
        params = {p.name: p.default for p in self.spec.contract.params}
        self.bound = bind.bind(self.spec, self.dataset, self.columns,
                               layout=layout, params=params)
        # findings from bind (e.g. the ordinal-coding note) BEFORE check overwrites
        self.bind_finding_codes = [f.code for f in self.bound.findings]
        profs = {p.name: p for p in self.dataset.profiles}
        # gate() mirrors the app seam: structural findings MERGED with bind's own
        # (dedup by code), so a bind-level block (T8's ORDORDER) is not overwritten
        # before it reaches the runner.
        self.bound.findings = check.gate(self.bound, profs)
        self.result = self.spec.run(self.bound)          # must not raise
        self.sentence = sentences.render(self.result)
        self.figs = charts.chart(self.result)

    def report_surfaces(self) -> tuple[list[str], list[str], str]:
        """(paragraphs, table-cells, joined-text) of a round-tripped .docx."""
        buf = report.build_report(self.result, self.spec,
                                  dataset=self.dataset, bound=self.bound)
        data = buf.getvalue()
        assert data[:2] == b"PK", "report is not a .docx (zip)"
        doc = Document(io.BytesIO(data))                 # raises if malformed
        paras = [p.text for p in doc.paragraphs]
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return paras, cells, "\n".join(paras + cells)


def _bare_float_label(text: str, digit_class: str):
    """Find a BARE 'N.0' float label (a leaked ``str(1.0)`` == '1.0'), while
    ignoring legitimate decimals like '-1.00', '13.50' or '21.05' that ``fmt`` /
    ``_f2`` produce. A real statistic value of 1.0 renders as '1' or '1.00',
    never '1.0', so a standalone 'N.0' can only be a leaked raw float label (F4).
    Mirrors the boundary-precise matchers in test_e2e.py."""
    return re.compile(r"(?<![\d.])[" + digit_class + r"]\.0(?![\d])").search(text)


# --------------------------------------------------------------------------
# Fixture (a): all-text survey -> NB1 + NB2 + S-H
#
# Group (Ctrl/Drug, 2 levels) + Rating (ordinal Likert TEXT). The first data row
# ('Ctrl','low') is deliberately short unit-ish tokens so it PASSES the units-row
# SHAPE gate -- a shape-only detector (pre-fix) would eat it. Only the shipped
# step-2 check ("a units row annotates NUMBERS") rejects it.
# --------------------------------------------------------------------------
_SURVEY = _csv([
    "Group,Rating",
    "Ctrl,low", "Ctrl,low", "Ctrl,medium", "Ctrl,low", "Ctrl,medium",
    "Ctrl,high", "Ctrl,low", "Ctrl,medium", "Ctrl,low", "Ctrl,high",
    "Drug,high", "Drug,medium", "Drug,high", "Drug,high", "Drug,medium",
    "Drug,high", "Drug,medium", "Drug,high", "Drug,high", "Drug,medium",
])


def test_all_text_survey_nb1_first_row_not_eaten_as_units():
    """NB1: an all-text sheet keeps its first data row. Pre-fix ``_is_units_row``
    matched only on token SHAPE, so 'Ctrl','low' was consumed as a units row ->
    n one short (19), units set, and Excel row 2 logged as dropped -- silently."""
    t = clean.clean(grid.load(_SURVEY, "survey.csv")[0])
    assert len(t.rows) == 20                       # pre-fix: 19 (first row eaten)
    assert t.units_row is None                     # pre-fix: 2 (the units row)
    assert all(u is None for u in t.units)         # pre-fix: ['Ctrl','low']
    dropped = {e for e, _r in t.dropped_rows}
    assert 2 not in dropped                         # Excel row 2 = first data row
    assert t.excel_rows[0] == 2 and t.excel_rows[-1] == 21


def test_all_text_survey_sh_not_header_ambiguous():
    """S-H: a clean all-text survey is NOT flagged header-ambiguous. Pre-fix every
    row of an all-text sheet scored ~1.0 (all-text, high strfrac), so a data row
    tied the header and the sheet forced the header picker open."""
    t = clean.clean(grid.load(_SURVEY, "survey.csv")[0])
    assert t.ambiguous is False                    # pre-fix: True (rivals tie)
    assert t.header == ["Group", "Rating"]         # header row, not a data row
    assert t.header_row == 1


def test_all_text_survey_nb2_ordinal_rank_test_does_not_crash():
    """NB2: an ordinal Likert TEXT outcome runs a rank test cleanly. Pre-fix the
    string labels ('low'/'medium'/'high') reached ``to_numpy(float)`` inside mwu
    and raised ``ValueError: could not convert string to float``."""
    r = _Run(_SURVEY, "survey.csv", "mwu", "long")   # would RAISE pre-fix
    # the coding choke turned the ordinal TEXT column into ranks before the runner:
    assert "ORDCODE" in r.bind_finding_codes
    assert pd.api.types.is_numeric_dtype(r.bound.data["outcome"])
    assert set(r.bound.data["outcome"].dropna().tolist()) == {1.0, 2.0, 3.0}
    # a real Result (never an uncaught crash) with a non-empty plain-English line
    assert r.result.status == "ok"
    assert r.sentence.strip()
    assert r.result.groups == ("Ctrl", "Drug")


# --------------------------------------------------------------------------
# Fixture (b): numeric-coded long-paired -> S-B(a) / F4
#
# Subject / Time (coded 1,2) / Score. The long->wide pivot derives the condition
# labels; they must read "1"/"2", not the float "1.0"/"2.0", everywhere.
# --------------------------------------------------------------------------
def _paired_csv() -> bytes:
    # Score deliberately carries duplicate values (58/60 recur) so it is NOT all-
    # distinct: an all-distinct numeric column with m>=20 infers as ID, not NUMERIC
    # (infer._numeric_kinds), and then no column could fill the numeric outcome.
    rows = ["Subject,Time,Score"]
    before = [42, 48, 50, 55, 47, 50, 44, 48, 46, 52, 49, 45]
    after = [58, 60, 66, 60, 60, 64, 58, 68, 60, 63, 62, 58]
    for i in range(12):
        rows.append(f"S{i + 1},1,{before[i]}")
        rows.append(f"S{i + 1},2,{after[i]}")
    return _csv(rows)


_PAIRED = _paired_csv()


@pytest.mark.parametrize("test_id", ["t_paired", "wilcoxon"])
def test_numeric_coded_condition_labels_are_integers(test_id):
    """S-B(a)/F4: after long->wide, the derived condition labels render as the
    integers "1"/"2" in the Result labels, the plain-English sentence, and the
    Word report. Pre-fix ``_long_to_wide`` derived them via ``str(float)`` -> the
    labels, sentence and report all read "1.0"/"2.0"."""
    r = _Run(_PAIRED, "paired.csv", test_id, "long")
    assert r.result.status == "ok"

    # Result labels: the exact fixed output (pre-fix: "1.0" / "2.0").
    assert r.result.labels["before"] == "1"
    assert r.result.labels["after"] == "2"
    assert str(r.result.higher) in ("1", "2")      # a label, not a raw float

    # the sentence names the conditions as integers, never as floats.
    assert "between 1 and 2" in r.sentence          # pre-fix: "between 1.0 and 2.0"
    assert _bare_float_label(r.sentence, "12") is None, \
        f"a float condition label leaked into the sentence: {r.sentence!r}"

    # the Word report the student downloads. Scan the TABLE CELLS (where the
    # labels live), not the whole doc -- the Methods prose carries library
    # version numbers like "pandas 3.0.6" that are not data labels.
    _paras, cells, _rtext = r.report_surfaces()
    assert "1" in cells and "2" in cells            # provenance rows: '1' / '2'
    assert _bare_float_label("\n".join(cells), "12") is None, \
        "a float condition label ('1.0'/'2.0') leaked into the report table"


# --------------------------------------------------------------------------
# Fixture (c): die-face goodness-of-fit -> S-B(b) / F4
#
# A raw column of die faces coded 1..6 (chi2_gof binds it as the 'category' role
# in the long layout). The face labels must read "1".."6", never "1.0".."6.0".
# --------------------------------------------------------------------------
def _die_csv() -> bytes:
    rows = ["Face"]
    for face in range(1, 7):
        rows += [str(face)] * (10 + face)          # faces 1..6, counts 11..16
    return _csv(rows)


_DIE = _die_csv()


def test_die_face_gof_labels_are_integers():
    """S-B(b)/F4: chi2_gof over numeric-coded die faces renders the category
    labels as "1".."6" in the Result groups/table and the Word report. Pre-fix
    ``cat.py`` rendered the raw floats -> "1.0".."6.0"."""
    r = _Run(_DIE, "die.csv", "chi2_gof", "long")
    assert r.result.status == "ok"

    # Result groups: the exact fixed output (pre-fix: "1.0".."6.0").
    assert r.result.groups == ("1", "2", "3", "4", "5", "6")
    assert list(r.result.table["category"]) == ["1", "2", "3", "4", "5", "6"]
    assert r.sentence.strip()

    # the Word report table shows the six faces as integers, no "N.0" leak.
    # Scan the TABLE CELLS only (the Methods prose carries library versions).
    _paras, cells, _rtext = r.report_surfaces()
    for face in ("1", "2", "3", "4", "5", "6"):
        assert face in cells, f"face {face} missing from the report table"
    assert _bare_float_label("\n".join(cells), "1-6") is None, \
        "a float face label ('1.0'..'6.0') leaked into the report table"


# ==========================================================================
# Wave #3 -- the corpus-gap SEAM defects (see the module docstring's table).
# Every numeric GREEN reference is computed IN the test from scipy over hand-
# built Python lists (never from the pipeline's own output). Crash pins call
# _Run directly: the RED is the propagated exception, the GREEN the same call
# succeeding -- no pytest.raises, no try/except.
# ==========================================================================

# T1/T2: a survey with a blank header cell. The 24 data rows cycle four 2-level
# columns on different periods, so exactly one recurs to a repeated-header drop
# and the first data row reads M/Yes/North/Placebo (the header current clean
# steals). Header names differ (T1 trailing blank -> 'Column D'; T2 leading
# blank -> 'Column A'); the body is identical.
_SURVEY_ROWS = [
    f"{['M', 'F'][i % 2]},{['Yes', 'No'][(i // 2) % 2]},"
    f"{['North', 'South'][(i // 4) % 2]},{['Placebo', 'Drug'][(i // 3) % 2]}"
    for i in range(24)
]
_T1 = _csv(["Sex,Smoker,Region,"] + _SURVEY_ROWS)
_T2 = _csv([",Smoker,Region,Arm"] + _SURVEY_ROWS)


def test_t1_trailing_blank_header_cell_not_eaten():
    """NB3: a header whose last cell is blank must still win over the first data
    row. Pre-fix the blank lowered the header's fill so the all-text data row
    'M,Yes,North,Placebo' scored higher and became the header (header_row==2,
    22 rows), and chi2_ind then bound row='M'/col='Yes'."""
    t = clean.clean(grid.load(_T1, "t1.csv")[0])
    assert t.header_row == 1                        # pre-fix: 2 (first data row)
    assert t.header == ["Sex", "Smoker", "Region", "Column D"]
    assert t.ambiguous is False
    assert len(t.rows) == 24                         # pre-fix: 22
    assert t.excel_rows[0] == 2
    r = _Run(_T1, "t1.csv", "chi2_ind", "long")
    assert r.columns["row"] == ("Sex",) and r.columns["col"] == ("Smoker",)
    assert r.result.n["total"] == 24                 # pre-fix: 22


def test_t2_leading_blank_header_cell_not_eaten():
    """NB3 (mirror): a header whose FIRST cell is blank is still the header.
    Pre-fix header_row==2, the first data row was eaten (22 rows)."""
    t = clean.clean(grid.load(_T2, "t2.csv")[0])
    assert t.header_row == 1                         # pre-fix: 2
    assert t.header == ["Column A", "Smoker", "Region", "Arm"]
    assert t.ambiguous is False
    assert len(t.rows) == 24                         # pre-fix: 22
    assert t.excel_rows[0] == 2


# T3: sub-1 doses. Score is plain numeric; Dose is a mixed numeric+text column
# (three tiny doses + 'control'), whose CATEGORICAL df labels must stay lossless.
_T3 = _csv(["Dose,Score"] + [
    f"{['0.001', '0.002', '0.005', 'control'][i % 4]},{20 + (i % 7) * 1.3:.1f}"
    for i in range(24)
])


def test_t3_subunit_doses_do_not_merge_into_one_group():
    """S-D/F4: a mixed numeric+text column's CATEGORICAL labels must be lossless.
    Pre-fix the df column was built via ``fmt.num`` (2 dp), so 0.001 and 0.002
    both rendered '0' and MERGED into one ANOVA group -- although the PROFILE
    still (correctly) counted four levels. groups==('0','0.01','control'),
    n['0']==12: two distinct doses silently pooled."""
    ds = infer.infer(clean.clean(grid.load(_T3, "dose.csv")[0]))
    prof = next(p for p in ds.profiles if p.name == "Dose")
    assert prof.n_levels == 4                        # the profile always saw 4
    r = _Run(_T3, "dose.csv", "anova_1w", "long")
    assert r.result.status == "ok"
    # pre-fix: ('0', '0.01', 'control') -- 0.001 & 0.002 collapsed onto '0'.
    assert r.result.groups == ("0.001", "0.002", "0.005", "control")
    assert all(r.result.n[g] == 6
               for g in ("0.001", "0.002", "0.005", "control"))


# T4: a mixed numeric+junk column read for a t-test. The 20 real values are 3 dp;
# five junk-text A rows drop. The runner must use the RAW values, not 2-dp cells.
_T4_A = [0.012, 0.015, 0.013, 0.018, 0.014, 0.016, 0.011, 0.017, 0.019, 0.012]
_T4_B = [0.028, 0.031, 0.033, 0.035, 0.029, 0.030, 0.034, 0.032, 0.027, 0.031]
_T4 = _csv(["Grp,Conc"]
           + [f"A,{v}" for v in _T4_A] + [f"B,{v}" for v in _T4_B]
           + [f"A,{j}" for j in ["n/a?", "lost", "broke", "x", "?? "]])


def test_t4_mixed_column_ttest_uses_raw_not_rounded_values():
    """S-D/F4: a t-test over a mixed numeric+junk column must use the RAW cell
    values. Pre-fix the column was stored as 2-dp strings then re-parsed, so
    ``ttest_ind`` ran on the ROUNDED numbers: t==-8.878, p==2.4e-7. The true
    Welch t on the raw values is ~ -13.6626."""
    exp = ss.ttest_ind(_T4_A, _T4_B, equal_var=False)   # ~ -13.6626, 6.4e-11
    r = _Run(_T4, "t4.csv", "t_ind", "long")
    assert r.result.status == "ok"
    assert r.result.n["used"] == 20 and r.result.n["dropped"] == 5
    assert abs(r.result.statistic[1] - float(exp.statistic)) < 1e-9
    assert abs(r.result.p - float(exp.pvalue)) < 1e-9


# T5 / T13: a numeric column and a Yes/No column derived from it (Yes iff >=60).
_T5_SCORES = [55, 61, 48, 72, 66, 58, 80, 45, 70, 63,
              52, 77, 68, 49, 74, 60, 57, 83, 65, 71]
_T5 = _csv(["Score,Passed"]
           + [f"{s},{'Yes' if s >= 60 else 'No'}" for s in _T5_SCORES])


@pytest.mark.parametrize("test_id", ["pearson", "kendall", "ols_simple"])
def test_t5_binary_text_outcome_is_coded_not_crashed(test_id):
    """S-I: a Yes/No column used in a correlation/regression must be coded 0/1 at
    the bind choke. Pre-fix it reached ``to_numpy(float)`` uncoded and raised
    ``ValueError: could not convert string to float: 'No'`` inside spec.run (a
    crash pin: the RED is the propagated exception)."""
    passed01 = [1.0 if s >= 60 else 0.0 for s in _T5_SCORES]
    r = _Run(_T5, "t5.csv", test_id, "long")            # RED: ValueError 'No'
    # the binary column (y for pearson/kendall, x for ols_simple) is coded 0/1.
    coded = next(r.bound.data[role] for role, hdrs in r.bound.columns.items()
                 if "Passed" in hdrs)
    assert set(coded.dropna().tolist()) == {0.0, 1.0}
    assert "BINCODE" in r.bind_finding_codes
    bc = [f for f in r.result.findings if f.code == "BINCODE"]
    assert bc and all(tok in bc[0].text for tok in ("Passed", "Yes", "No"))
    if test_id == "pearson":
        exp = float(ss.pointbiserialr(passed01, _T5_SCORES).statistic)  # ~0.81267
        assert abs(r.result.statistic[1] - exp) < 1e-12
    elif test_id == "kendall":
        exp = float(ss.kendalltau(_T5_SCORES, passed01).statistic)      # ~0.69206
        assert abs(r.result.statistic[1] - exp) < 1e-12
    else:  # ols_simple: slope == mean(Yes) - mean(No) == 18.0
        yes = [s for s in _T5_SCORES if s >= 60]
        no = [s for s in _T5_SCORES if s < 60]
        exp = sum(yes) / len(yes) - sum(no) / len(no)                   # == 18.0
        assert exp == 18.0
        assert abs(r.result.extra["slope"] - exp) < 1e-9
    assert r.result.status == "ok"


# T6: a five-point satisfaction scale (a vocab the shipped ORDERED_VOCABS lacks).
_SAT = ["Very dissatisfied", "Dissatisfied", "Neutral", "Satisfied", "Very satisfied"]
_T6 = _csv(["Group,Satisfaction"]
           + [f"A,{_SAT[i % 5]}" for i in range(10)]
           + [f"B,{_SAT[(i + 2) % 5]}" for i in range(10)])


def test_t6_satisfaction_scale_is_ordinal():
    """S-J: the 5-point satisfaction scale is ORDINAL. Pre-fix it was not in any
    ORDERED_VOCAB, so kinds==(CATEGORICAL,) and mwu could not bind an outcome."""
    ds = infer.infer(clean.clean(grid.load(_T6, "sat.csv")[0]))
    prof = next(p for p in ds.profiles if p.name == "Satisfaction")
    assert prof.kinds[0] == Kind.ORDINAL            # pre-fix: CATEGORICAL
    assert prof.level_order == (
        "very dissatisfied", "dissatisfied", "neutral", "satisfied",
        "very satisfied")
    r = _Run(_T6, "sat.csv", "mwu", "long")
    assert "ORDCODE" in r.bind_finding_codes
    assert set(r.bound.data["outcome"].dropna().tolist()) <= {1.0, 2.0, 3.0, 4.0, 5.0}
    assert r.result.status == "ok"


# T7: low/medium/high WITH an explicit ORDINAL override (a known vocab column).
_T7 = _csv(["Group,Rating"]
           + [f"A,{['low', 'medium', 'high'][i % 3]}" for i in range(10)]
           + [f"B,{['low', 'medium', 'high'][(i + 1) % 3]}" for i in range(10)])


def test_t7_ordinal_override_on_known_vocab_keeps_order():
    """S-J: overriding a low/medium/high column to ORDINAL must keep its order.
    Pre-fix ORDINAL went down the numeric-family override path -> level_order
    None, levels (), a false '20 text cell(s) treated as missing' note, then
    ``ValueError: 'low'`` inside spec.run (a crash pin)."""
    r = _Run(_T7, "rating.csv", "mwu", "long",
             overrides={"Rating": Kind.ORDINAL})       # RED: ValueError 'low'
    prof = next(p for p in r.dataset.profiles if p.name == "Rating")
    assert prof.level_order == ("low", "medium", "high")
    assert prof.n_levels == 3
    assert prof.n_missing == 0
    assert not any("treated as missing" in n for n in prof.notes)
    assert "ORDCODE" in r.bind_finding_codes
    assert r.result.status == "ok"


# T8: an unknown 4-level scale WITH an ORDINAL override (no vocab to order it by).
_T8 = _csv(["Group,Rating"]
           + [f"A,{['Terrible', 'Bad', 'OK', 'Great'][i % 4]}" for i in range(10)]
           + [f"B,{['Terrible', 'Bad', 'OK', 'Great'][(i + 1) % 4]}" for i in range(10)])


def test_t8_ordinal_override_on_unknown_scale_blocks_not_crashes():
    """S-J: overriding an unknown scale to ORDINAL must BLOCK with a real message,
    not crash. Pre-fix the numeric-family override path raised a false 'treated as
    missing' note then ``ValueError: 'Terrible'`` inside spec.run (a crash pin)."""
    r = _Run(_T8, "unknown.csv", "mwu", "long",
             overrides={"Rating": Kind.ORDINAL})       # RED: ValueError 'Terrible'
    prof = next(p for p in r.dataset.profiles if p.name == "Rating")
    assert prof.n_levels == 4
    assert prof.n_missing == 0
    assert not any("treated as missing" in n for n in prof.notes)
    assert "ORDORDER" in r.bind_finding_codes
    assert r.result.status == "blocked"
    assert r.sentence.strip()
    # the block is ORDORDER itself, not a downstream S5; the column is KEPT (not
    # blanked) so every row survives and the message is about the order, never a
    # false "usable rows" / "missing" story. Pre-fix: findings[0]==S5, n_used==0.
    assert r.result.findings[0].code == "ORDORDER"
    assert r.bound.n_used == 20
    assert "order" in r.sentence and "Rating" in r.sentence
    assert "usable rows" not in r.sentence and "missing" not in r.sentence


# T9: a 22-row BMI column, 22 distinct 1-dp values (a real numeric measurement).
_T9_BMI = [21.3, 22.1, 23.4, 24.8, 25.2, 26.7, 27.1, 28.5, 29.9, 30.2, 19.8,
           20.4, 22.9, 23.7, 24.1, 25.9, 26.3, 27.8, 28.2, 29.4, 31.1, 32.6]
_T9 = _csv(["Group,BMI"]
           + [f"{'A' if i % 2 == 0 else 'B'},{v}" for i, v in enumerate(_T9_BMI)])


def test_t9_fractional_near_contiguous_column_is_numeric_not_id():
    """S-K(a): a 22-distinct 1-dp BMI column is NUMERIC (usable as an outcome).
    Pre-fix the near-contiguity ID rule ignored the fractional part and typed it
    ID -> kinds==(ID,) and t_ind reported 'needs a numeric column; your sheet has
    none'."""
    ds = infer.infer(clean.clean(grid.load(_T9, "bmi.csv")[0]))
    prof = next(p for p in ds.profiles if p.name == "BMI")
    assert prof.kinds == (Kind.NUMERIC, Kind.ID)    # pre-fix: (Kind.ID,)
    ok, _why = bind.satisfies(REGISTRY["t_ind"], ds)
    assert ok is True                                # pre-fix: False
    r = _Run(_T9, "bmi.csv", "t_ind", "long")
    assert r.result.status == "ok"


# T10: two integer near-contiguous columns (both typed ID by the same rule).
_T10_BEFORE = [40, 41, 43, 44, 46, 47, 49, 50, 52, 53,
               55, 56, 58, 59, 61, 62, 64, 65, 67, 68]
_T10 = _csv(["Before,After"]
            + [f"{v},{v + 5 + (i % 2)}" for i, v in enumerate(_T10_BEFORE)])


def test_t10_no_numeric_message_names_the_offending_columns():
    """S-K(b): when a paired test cannot bind because both columns are typed ID,
    the message must NAME them (not lie). Pre-fix satisfies(t_paired) returned the
    generic lie 'needs a numeric column; your sheet has none' -- though the sheet
    plainly holds two numeric columns. Pin only the four tokens."""
    ds = infer.infer(clean.clean(grid.load(_T10, "ba.csv")[0]))
    ok, why = bind.satisfies(REGISTRY["t_paired"], ds)
    assert ok is False
    assert "Before" in why and "After" in why       # pre-fix: names neither
    assert "your sheet has none" not in why          # pre-fix: present (the lie)
    assert "numeric" in why.casefold()


# T11: a small 2x3 count layout -> Fisher's Freeman-Halton (R x C) path.
_T11 = _csv(["Grp,Out", "A,x", "A,y", "A,z", "B,x", "B,x"])


def test_t11_fisher_rxc_does_not_crash_on_scalar_cast():
    """S-L: an R x C Fisher test must not crash. Pre-fix the Monte-Carlo p-value
    is an array and ``float(res.pvalue)`` raised ``TypeError: only 0-dimensional
    arrays can be converted to Python scalars`` at cat.py (a crash pin)."""
    r = _Run(_T11, "fisher.csv", "fisher", "long")      # RED: TypeError (0-d cast)
    assert r.result.status == "ok"
    assert 0.0 <= r.result.p <= 1.0
    assert r.result.statistic is not None and pd.notna(r.result.statistic[1])
    _paras, _cells, rtext = r.report_surfaces()          # docx round-trip
    assert r.sentence.strip() and rtext.strip()


# T12: a roster whose FIRST data row is all NA tokens (a real subject, missing).
_T12 = _csv(["Name,Height,Weight,Age", "Ali,NA,n/a,na",
             "Bob,170,65,30", "Cara,160,55,25", "Dan,180,80,40",
             "Eve,165,60,28", "Finn,175,70,35"])


def test_t12_na_token_data_row_is_not_a_units_row():
    """N13: a first data row of NA tokens is DATA (a subject with missing
    measurements), not a units row. Pre-fix it passed the units-row shape gate and
    was silently CONSUMED: units_row==2, units==['Ali','NA','n/a','na'], 5 rows,
    excel_rows[0]==3, nothing logged."""
    t = clean.clean(grid.load(_T12, "roster.csv")[0])
    assert t.units_row is None                       # pre-fix: 2
    assert len(t.rows) == 6                            # pre-fix: 5
    assert t.excel_rows[0] == 2                        # pre-fix: 3
    ds = infer.infer(t)
    prof = next(p for p in ds.profiles if p.name == "Height")
    assert prof.n_missing == 1                         # pre-fix: 0 (row was eaten)


def test_t13_partial_auto_bind_returns_none():
    """N12: ``auto_columns`` must return None when a required role has no eligible
    column, not a PARTIAL assignment. Pre-fix spearman (x,y both numeric/ordinal)
    on Score + Passed(binary) returned ``{'x': ('Score',), 'y': ()}`` -- an empty
    tuple for the unfillable y role instead of failing closed."""
    ds = infer.infer(clean.clean(grid.load(_T5, "s.csv")[0]))
    cols = bind.auto_columns(REGISTRY["spearman"], ds, "long")
    assert cols is None                              # pre-fix: {'x': ('Score',), 'y': ()}
