"""RED-first tests for statkit/report.py (Chunk 13, PLAN §7).

DoD slice: the docx ROUND-TRIP. ``build_report(result)`` assembles a python-docx
Word report entirely IN MEMORY and returns an ``io.BytesIO``; we reopen it with
``docx.Document(buf)`` and assert the report actually carries what §7 requires:

  * the plain-English sentence (``sentences.render``) appears as text;
  * a numbers table is present;
  * for a charted test, an inline image is embedded (chart -> BytesIO ->
    add_picture);
  * a short methods line is present;
  * a BLOCKED result states the block reason and embeds NO chart.

L3: the doc is built and saved to a BytesIO — never to disk (mechanical guard:
tests/test_no_ai_no_io.py; here the round-trip proves in-memory assembly).
"""
from __future__ import annotations

import io
import math

import numpy as np
import pandas as pd
import pytest
from docx import Document

from statkit import bind, sentences
from statkit.infer import Dataset
from statkit.model import Check, ColumnProfile, Kind
from statkit.registry import REGISTRY

N, O, CAT, B, ID = (Kind.NUMERIC, Kind.ORDINAL, Kind.CATEGORICAL, Kind.BINARY, Kind.ID)


# --------------------------------------------------------------------------
# compact dataset / bound builders (same shape as tests/test_means.py)
# --------------------------------------------------------------------------
def _isna(v):
    return v is None or v is pd.NA or (isinstance(v, float) and math.isnan(v))


def _profile(name, kinds, values):
    non_null = [v for v in values if not _isna(v)]
    if kinds[0] in (Kind.NUMERIC, Kind.ID, Kind.DATE, Kind.EMPTY):
        levels = ()
    else:
        seen, levels = set(), []
        for v in non_null:
            s = str(v)
            if s not in seen:
                seen.add(s); levels.append(s)
        levels = tuple(levels)
    return ColumnProfile(
        name=name, kinds=tuple(kinds), n_total=len(values),
        n_missing=len(values) - len(non_null), n_levels=len(levels), levels=levels)


def _ds(cols):
    df_cols, profiles, n = {}, [], None
    for name, (values, kinds) in cols.items():
        n = len(values) if n is None else n
        non_null = [v for v in values if not _isna(v)]
        numeric = all(isinstance(v, (int, float)) and not isinstance(v, bool)
                      for v in non_null)
        if kinds[0] in (N, O, B, ID) and numeric:
            df_cols[name] = pd.array(
                [None if _isna(v) else float(v) for v in values], dtype="float64")
        else:
            df_cols[name] = pd.array(
                [pd.NA if _isna(v) else str(v) for v in values], dtype="string")
        profiles.append(_profile(name, kinds, values))
    return Dataset(df=pd.DataFrame(df_cols), profiles=tuple(profiles),
                   excel_rows=tuple(range(2, 2 + (n or 0))))


def _run(test_id, cols, columns, params=None):
    b = bind.bind(REGISTRY[test_id], _ds(cols), columns, params=params)
    return REGISTRY[test_id].run(b)


_R = np.random.default_rng(7)


def _t_ind():
    return _run("t_ind",
                {"Y": (list(np.round(_R.normal(0, 1, 12), 2))
                       + list(np.round(_R.normal(2, 1, 12), 2)), (N,)),
                 "G": (["A"] * 12 + ["B"] * 12, (CAT,))},
                {"outcome": ("Y",), "group": ("G",)})


def _pearson():
    x = list(np.round(_R.normal(10, 3, 20), 2))
    y = list(np.round(np.array(x) * 1.4 + _R.normal(0, 2, 20), 2))
    return _run("pearson", {"X": (x, (N,)), "Y": (y, (N,))},
                {"x": ("X",), "y": ("Y",)})


def _chi2():
    row = ["A"] * 20 + ["B"] * 20
    col = (["X"] * 15 + ["Y"] * 5) + (["X"] * 5 + ["Y"] * 15)
    return _run("chi2_ind", {"R": (row, (CAT,)), "Cc": (col, (CAT,))},
                {"row": ("R",), "col": ("Cc",)})


def _ols():
    x = list(np.round(_R.normal(10, 3, 20), 2))
    y = list(np.round(np.array(x) * 1.4 + _R.normal(0, 2, 20), 2))
    return _run("ols_simple", {"Y": (y, (N,)), "X": (x, (N,))},
                {"outcome": ("Y",), "x": ("X",)})


def _describe():
    return _run("describe", {"Y": ([1, 2, 3, 4, 5, 6, 7, 8], (N,))},
                {"variables": ("Y",)})


def _blocked():
    return _run("t_ind", {"Y": ([5, 5, 5, 5, 5, 5], (N,)),
                          "G": (["A"] * 3 + ["B"] * 3, (CAT,))},
                {"outcome": ("Y",), "group": ("G",)})


# --------------------------------------------------------------------------
# helpers on a reopened Document
# --------------------------------------------------------------------------
def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _reopen(buf):
    buf.seek(0)
    return Document(buf)


def _find_table(doc, header):
    """The first table whose header row contains ``header``, or None."""
    for t in doc.tables:
        if header in [c.text for c in t.rows[0].cells]:
            return t
    return None


def _headings(doc):
    return [p.text for p in doc.paragraphs
            if p.style is not None and p.style.name.startswith("Heading")]


# ==========================================================================
# docx round-trip: BytesIO in, reopened Document out
# ==========================================================================
def test_build_report_returns_bytesio_docx():
    from statkit import report
    buf = report.build_report(_t_ind(), spec=REGISTRY["t_ind"])
    assert isinstance(buf, io.BytesIO)
    assert buf.getvalue()[:2] == b"PK"          # docx is a zip
    _reopen(buf)                                 # reopens without error


@pytest.mark.parametrize("factory,test_id,charted", [
    (_t_ind, "t_ind", True),
    (_pearson, "pearson", True),
    (_chi2, "chi2_ind", True),
    (_ols, "ols_simple", True),
    (_describe, "describe", False),
])
def test_report_round_trip_per_family(factory, test_id, charted):
    from statkit import report
    r = factory()
    assert r.status == "ok", [f.text for f in r.findings]
    buf = report.build_report(r, spec=REGISTRY[test_id])
    doc = _reopen(buf)

    # 1) the plain-English sentence is present verbatim
    sentence = sentences.render(r)
    body = _text(doc)
    assert sentence[:40] in body, test_id

    # 2) a numbers table is present
    assert len(doc.tables) >= 1, test_id

    # 3) a charted test embeds an inline image; a chart-less one does not
    n_images = len(doc.inline_shapes)
    if charted:
        assert n_images >= 1, test_id
    else:
        assert n_images == 0, test_id

    # 4) a methods line is present
    assert "Missing values were excluded" in body or "Analysis used" in body, test_id


def test_blocked_report_states_reason_and_has_no_chart():
    from statkit import report
    r = _blocked()
    assert r.status == "blocked"
    buf = report.build_report(r, spec=REGISTRY["t_ind"])
    doc = _reopen(buf)
    body = _text(doc)
    # the sentence IS the block reason (PLAN §6 rule 18)
    assert sentences.render(r)[:30] in body
    # no chart on a blocked result
    assert len(doc.inline_shapes) == 0


def test_report_embeds_two_figures_for_ols_simple():
    # ols_simple = scatter_fit + resid_panels (§7) -> two inline images
    from statkit import report
    doc = _reopen(report.build_report(_ols(), spec=REGISTRY["ols_simple"]))
    assert len(doc.inline_shapes) == 2


def test_report_without_spec_still_builds():
    # spec is optional; the report degrades gracefully (no citation line).
    from statkit import report
    buf = report.build_report(_t_ind())
    assert _reopen(buf) is not None


# ==========================================================================
# Batch 1G — B2 (APA p-cells), S11 (Data section), S4 render, NICE ω²
# ==========================================================================
def test_table_p_values_are_apa_not_zero():
    # B2: a real post-hoc p of .0029 must render as ".003" (APA), never "0"
    # (fmt.num(.0029)=="0"); and a plain RangeIndex must NOT print a bare 0/1/2
    # index column — the first header cell is the first real column.
    from statkit import report
    from statkit.model import Result
    ph = pd.DataFrame({"group1": ["A"], "group2": ["B"],
                       "p": [0.0029], "p_holm": [0.0029]})
    r = Result(test_id="posthoc_fixture", test_name="Post-hoc demo",
               status="ok", posthoc=ph, posthoc_name="Dunn")
    doc = _reopen(report.build_report(r))
    tbl = _find_table(doc, "group1")
    assert tbl is not None, "post-hoc table not rendered"
    assert tbl.rows[0].cells[0].text == "group1"        # no bare index column
    data = [c.text for row in tbl.rows[1:] for c in row.cells]
    assert any(t in ("p = .003", ".003") for t in data), data
    assert "0" not in data, data


def test_data_section_lists_columns_prep_log_and_n_accounting():
    # S11: dataset/bound populate a "Data" section — columns-used table (role
    # label + header + kind), the prep log, N accounting, and Appendix A rows.
    from types import SimpleNamespace
    from statkit import report
    r = _t_ind()
    spec = REGISTRY["t_ind"]
    dataset = SimpleNamespace(
        source="grades.xlsx", sheet="Fall", n_rows_read=24, header_row=1,
        log=("Detected header on row 1",
             "Dropped 1 summary row 'Class Average'"),
        dropped_rows=((30, "clean: summary row 'Class Average'"),),
        profiles=(),
    )
    bound = SimpleNamespace(
        layout="long",
        columns={"outcome": ("Score",), "group": ("Class",)},
        kinds={"outcome": Kind.NUMERIC, "group": Kind.CATEGORICAL},
        n_total=25, n_used=24, dropped={"missing outcome": 1},
        dropped_rows=((13, "missing outcome"),),
    )
    doc = _reopen(report.build_report(r, spec=spec, dataset=dataset, bound=bound))
    headings = _headings(doc)
    body = _text(doc)
    assert "Data" in headings
    cols_tbl = _find_table(doc, "Role")
    assert cols_tbl is not None, "columns-used table missing"
    cell_text = [c.text for row in cols_tbl.rows for c in row.cells]
    assert "Score" in cell_text and "Class" in cell_text
    assert "Outcome (numeric)" in cell_text
    assert "Group (2 levels)" in cell_text
    assert "Detected header on row 1" in body          # prep-log line
    assert any(h.startswith("Appendix A") for h in headings)
    appA = _find_table(doc, "Excel row")
    assert appA is not None
    assert "13" in [c.text for row in appA.rows for c in row.cells]
    assert "missing outcome" in body                   # N-accounting reason


def test_known_extra_tables_are_rendered_with_captions():
    # S4: a chi-square Result carries extra["observed"] -> "Observed counts".
    from statkit import report
    r = _chi2()
    assert r.status == "ok"
    assert isinstance(r.extra.get("observed"), pd.DataFrame)
    doc = _reopen(report.build_report(r, spec=REGISTRY["chi2_ind"]))
    assert "Observed counts" in _text(doc)


def test_or_row_says_not_estimable_for_empty_cell():
    # F2: an infinite odds ratio (empty 2×2 cell) must read "not estimable
    # (empty cell)" in the numbers panel, never a bare "inf"/"—" (+ no CI).
    from statkit import report
    from statkit.model import Result
    r = Result(test_id="fisher", test_name="Fisher's exact test", status="ok",
               p=0.5, estimate=("odds ratio", float("inf")),
               estimate_ci=(1.5, float("inf")),
               n={"total": 10, "used": 10, "dropped": 0})
    doc = _reopen(report.build_report(r))
    body = _text(doc)
    assert "not estimable (empty cell)" in body


def test_describe_categorical_only_renders_counts_table():
    # S-C: describe run on ONLY categorical columns has descriptives is None, so
    # today the report shows NO results table at all — the student gets nothing,
    # even though the summary sentence promises "counts ... (categorical)". The
    # report MUST render extra["categoricals"] as a counts table.
    from statkit import report
    r = _run("describe",
             {"Sex": (["M", "F", "F", "M", "F", None], (CAT,))},
             {"variables": ("Sex",)})
    assert r.status == "ok"
    assert r.descriptives is None
    assert r.extra.get("categoricals") == {"Sex": {"F": 3, "M": 2}}
    doc = _reopen(report.build_report(r, spec=REGISTRY["describe"]))
    tbl = _find_table(doc, "Sex")
    assert tbl is not None, "categorical counts table missing"
    cells = [c.text for row in tbl.rows for c in row.cells]
    assert "M" in cells and "F" in cells        # the level labels
    assert "3" in cells and "2" in cells         # their counts (F=3, M=2)


def test_describe_grouped_categorical_renders_level_by_group_counts():
    # S-C (grouped shape): {header: {group: {level: count}}} -> a level × group
    # count matrix (group labels as columns, levels as rows).
    from statkit import report
    r = _run("describe",
             {"Sex": (["M", "F", "F", "M", "F", "M"], (CAT,)),
              "Grp": (["A", "B", "A", "B", "A", "B"], (CAT,))},
             {"variables": ("Sex",), "group": ("Grp",)})
    assert r.status == "ok"
    assert r.descriptives is None
    cats = r.extra.get("categoricals")
    assert cats and all(isinstance(v, dict) for v in cats["Sex"].values())
    doc = _reopen(report.build_report(r, spec=REGISTRY["describe"]))
    tbl = _find_table(doc, "Sex")
    assert tbl is not None, "grouped categorical counts table missing"
    header = [c.text for c in tbl.rows[0].cells]
    assert "A" in header and "B" in header       # group labels as columns
    cells = [c.text for row in tbl.rows for c in row.cells]
    assert "M" in cells and "F" in cells          # level labels as rows


def _perfect_fit_ols():
    from statkit.model import Result
    table = pd.DataFrame(
        {"coef": [2.0, 1.5], "ci_low": [1.5, 1.5], "ci_high": [1.5, 1.5],
         "p": [0.0, 0.0]}, index=["Intercept", "X"])
    return Result(
        test_id="ols_simple", test_name="Simple linear regression", status="ok",
        statistic=("F", 1.5449491690281546e31), df=(1.0, 12.0), p=0.0,
        estimate=("slope", 1.5), estimate_ci=(1.5, 1.5),
        effect=("f2", float("inf")), effect_label="large",
        effect_source="Cohen (1988)",
        n={"total": 14, "used": 14, "dropped": 0}, table=table,
        labels={"outcome": "Y", "x": "X"},
        extra={"intercept": 2.0, "slope": 1.5, "r2": 1.0, "adj_r2": 1.0})


def test_perfect_fit_report_has_no_inf_no_astronomical_f_and_a_note():
    # S-E: a perfect-fit OLS (Cohen's f² = inf, F astronomically large) must NOT
    # print "inf" or a 30-digit F in the docx; instead a "perfect fit" note.
    import re
    from statkit import report
    body = _text(_reopen(report.build_report(
        _perfect_fit_ols(), spec=REGISTRY["ols_simple"])))
    assert "inf" not in body.lower(), body
    assert not re.search(r"\d{12,}", body), "astronomical F leaked into report"
    assert "perfect fit" in body.lower()


def test_finite_regression_report_still_prints_f_to_two_decimals():
    # Guard: a normal (finite) regression is unchanged — F still renders ".2f".
    import re
    from statkit import report
    body = _text(_reopen(report.build_report(_ols(), spec=REGISTRY["ols_simple"])))
    assert "perfect fit" not in body.lower()
    # the F statistic still shows two decimals (a known finite value ends ".NN")
    assert re.search(r"F\(1, 18\) = \d+\.\d\d", body), body


def test_assumption_table_never_prints_nan_or_inf():
    # S-O: a non-finite statistic/p in a Check must never reach the docx as
    # "nan"/"inf"; the statistic cell uses fmt.num (-> em dash), p already guarded.
    from statkit import report
    r = _t_ind()
    r.checks = (Check("Brown-Forsythe (Levene, center=median)", math.nan, math.nan, False, ""),
                Check("Bartlett", math.inf, 0.0, False, ""))
    doc = _reopen(report.build_report(r, REGISTRY["t_ind"]))
    tbl = _find_table(doc, "Check")
    cells = [c.text for row in tbl.rows for c in row.cells]
    assert not any(c.strip().lower() in ("nan", "inf", "-inf") for c in cells), cells
    assert "—" in cells                       # fmt.NON_FINITE for the inf statistic
    # the real path: a two-valued outcome per group -> Brown-Forsythe p is nan
    r2 = _run("t_ind", {"Y": ([0, 1] * 12, (N,)), "G": (["A"] * 12 + ["B"] * 12, (CAT,))},
              {"outcome": ("Y",), "group": ("G",)})
    assert r2.checks[0].name.startswith("Brown-Forsythe")
    assert r2.checks[0].passed is None
    assert r2.checks[0].p is None
    doc2 = _reopen(report.build_report(r2, REGISTRY["t_ind"]))
    for t in doc2.tables:
        for row in t.rows:
            for c in row.cells:
                assert c.text != "nan", "nan leaked into a report cell"


def test_omega_row_present_for_anova():
    # NICE: an ANOVA Result with extra["omega_sq"] gets an "ω²" APA row.
    from statkit import report
    y = (list(np.round(_R.normal(0, 1, 10), 2))
         + list(np.round(_R.normal(1, 1, 10), 2))
         + list(np.round(_R.normal(2, 1, 10), 2)))
    g = ["A"] * 10 + ["B"] * 10 + ["C"] * 10
    r = _run("anova_1w", {"Y": (y, (N,)), "G": (g, (CAT,))},
             {"outcome": ("Y",), "group": ("G",)})
    assert r.status == "ok"
    assert r.extra.get("omega_sq") is not None
    doc = _reopen(report.build_report(r, spec=REGISTRY["anova_1w"]))
    assert "ω²" in _text(doc)
