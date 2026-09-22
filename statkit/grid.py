"""bytes -> Grid: the raw spreadsheet reader (PLAN §4.1).

The first stage of the ingest pipeline. It turns an uploaded file (as bytes,
never a disk path -- L3: Streamlit hands us an in-memory buffer) into one
``Grid`` per non-empty sheet, preserving exactly what later stages need:

  * the raw cell value,
  * the number-format string ``fmt`` (coerce reads it for percent detection),
  * whether a cell was a cached formula vs a literal,
  * merged-cell ranges (1-based, inclusive),
  * hidden rows / columns (1-based),
  * the 1-based Excel row numbers, preserved end-to-end through the pipeline.

Reading is buffer-only: openpyxl/xlrd on ``BytesIO`` / ``file_contents``,
csv on a decoded ``StringIO``. No builtin ``open`` on a path, no disk writes --
the L3 gate (tests/test_no_ai_no_io.py) enforces this over the whole package.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO

import openpyxl
import xlrd
from xlrd.xldate import xldate_as_datetime

# Sentinel for a formula cell whose cached value is absent (openpyxl reads an
# uncomputed formula as None under data_only=True). infer.py blocks such a
# column: "open in Excel, press Save, re-upload".
UNCOMPUTED = object()


class GridError(ValueError):
    """A file could not be read as the workbook its extension claims (S14).

    Raised with a message aimed at the student uploading it -- what went wrong
    and the concrete Excel step (Save As) that fixes it -- not a raw parser
    traceback.
    """

# Above this many cells a sheet is read in openpyxl's read_only fast mode:
# values only, no number formats / merged ranges / hidden flags (PLAN §4.1).
LARGE_SHEET_CELLS = 400_000

_LARGE_NOTE = (
    "large sheet (>400,000 cells): read in fast mode; "
    "number formats, merged cells and hidden rows/cols were not read"
)


@dataclass
class Cell:
    value: object
    fmt: str = "General"
    formula: bool = False


@dataclass
class Grid:
    rows: list  # list[list[Cell]] -- rectangular
    excel_rows: list  # list[int], 1-based sheet row numbers, one per row
    merged: list  # list[tuple[int, int, int, int]] -- (r1, c1, r2, c2) 1-based inclusive
    hidden_rows: frozenset  # frozenset[int], 1-based
    hidden_cols: frozenset  # frozenset[int], 1-based
    sheet: str
    source: str
    note: str = ""


def load(data: bytes, filename: str) -> list[Grid]:
    """Read a spreadsheet from bytes into one Grid per non-empty sheet.

    Dispatch is by ``filename`` extension: xlsx/xlsm -> openpyxl, xls -> xlrd,
    anything else -> delimited text (csv/txt). Sheets are returned in workbook
    order; a sheet with no data is skipped.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("xlsx", "xlsm"):
        return _read_workbook(_load_xlsx, data, filename, ext)
    if ext == "xls":
        return _read_workbook(_load_xls, data, filename, ext)
    return _load_csv(data, filename)


def _looks_like_text(data: bytes) -> bool:
    """A file that opens as plain UTF-8 text (or starts with '<') is not a real
    binary workbook -- almost always a CSV/HTML export renamed to .xlsx/.xls."""
    if data.lstrip()[:1] == b"<":
        return True
    try:
        data.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def _read_workbook(reader, data: bytes, filename: str, ext: str) -> list[Grid]:
    """Run a workbook reader, turning any parser failure into a GridError whose
    message tells the student what to do (S14). We branch on the bytes, never on
    an imported zip/parser type, to keep the L3 no-io gate green."""
    try:
        return reader(data, filename)
    except GridError:
        raise
    except Exception as e:
        if _looks_like_text(data):
            raise GridError(
                f"'{filename}' looks like a text or HTML file renamed .{ext}, "
                f"not a real {ext} workbook. Open it in Excel and use Save As to "
                f"save a genuine .{ext} file, then re-upload."
            ) from e
        raise GridError(
            f"'{filename}' is not a readable {ext} workbook "
            f"({type(e).__name__}). If it was exported from another program, "
            f"open it in Excel and use Save As."
        ) from e


# --------------------------------------------------------------------------
# xlsx / xlsm
# --------------------------------------------------------------------------
def _load_xlsx(data: bytes, filename: str) -> list[Grid]:
    # Cheap dimension peek first (read_only reads the <dimension> tag, not the
    # cells) so a huge sheet never gets fully loaded into memory before we
    # notice its size -- that is the whole point of the fallback.
    peek = openpyxl.load_workbook(BytesIO(data), read_only=True)
    try:
        big = any(
            (ws.max_row or 0) * (ws.max_column or 0) > LARGE_SHEET_CELLS
            for ws in peek.worksheets
        )
    finally:
        peek.close()

    if big:
        # ponytail: if ANY sheet is huge, read the whole workbook values-only.
        # Small sibling sheets in the same book lose fmt/merged/hidden, but a
        # huge sheet mixed with small ones is exotic; the realistic large case
        # is a single-sheet export. Upgrade to per-sheet routing if it matters.
        return _load_xlsx_readonly(data, filename)
    return _load_xlsx_full(data, filename)


def _load_xlsx_full(data: bytes, filename: str) -> list[Grid]:
    wb = openpyxl.load_workbook(BytesIO(data), data_only=False)
    wb_cached = None  # second load, only if some sheet has formulas
    grids: list[Grid] = []
    for ws in wb.worksheets:
        rows: list[list[Cell]] = []
        has_formula = False
        for row in ws.iter_rows():
            cells = []
            for c in row:
                is_f = c.data_type == "f"
                has_formula = has_formula or is_f
                cells.append(Cell(c.value, c.number_format or "General", is_f))
            rows.append(cells)

        if has_formula:
            if wb_cached is None:
                wb_cached = openpyxl.load_workbook(BytesIO(data), data_only=True)
            wsd = wb_cached[ws.title]
            for r_i, rowd in enumerate(wsd.iter_rows()):
                for c_i, cd in enumerate(rowd):
                    cell = rows[r_i][c_i]
                    if cell.formula:
                        cell.value = cd.value if cd.value is not None else UNCOMPUTED

        merged = [
            (m.min_row, m.min_col, m.max_row, m.max_col)
            for m in ws.merged_cells.ranges
        ]
        hidden_rows = frozenset(i for i, d in ws.row_dimensions.items() if d.hidden)
        hidden_cols = _hidden_cols(ws)

        rows, excel_rows = _finalize(rows)
        if not rows:
            continue  # empty sheet -> skip
        grids.append(
            Grid(rows, excel_rows, merged, hidden_rows, hidden_cols, ws.title, filename)
        )
    return grids


def _load_xlsx_readonly(data: bytes, filename: str) -> list[Grid]:
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        grids: list[Grid] = []
        for ws in wb.worksheets:
            rows = [
                [Cell(v, "General", False) for v in row]
                for row in ws.iter_rows(values_only=True)
            ]
            rows, excel_rows = _finalize(rows)
            if not rows:
                continue
            grids.append(
                Grid(rows, excel_rows, [], frozenset(), frozenset(),
                     ws.title, filename, note=_LARGE_NOTE)
            )
        return grids
    finally:
        wb.close()


def _hidden_cols(ws) -> frozenset:
    out: set = set()
    for d in ws.column_dimensions.values():
        if d.hidden and d.min:
            out.update(range(d.min, (d.max or d.min) + 1))
    return frozenset(out)


# --------------------------------------------------------------------------
# xls (legacy) via xlrd
# --------------------------------------------------------------------------
def _load_xls(data: bytes, filename: str) -> list[Grid]:
    wb = xlrd.open_workbook(file_contents=data, formatting_info=True)
    grids: list[Grid] = []
    for sh in wb.sheets():
        rows = [
            [_xls_cell(sh.cell(r, c), wb.datemode) for c in range(sh.ncols)]
            for r in range(sh.nrows)
        ]
        # xlrd merged_cells: (rlo, rhi, clo, chi), rhi/chi half-open 0-based.
        merged = [
            (rlo + 1, clo + 1, rhi, chi) for (rlo, rhi, clo, chi) in sh.merged_cells
        ]
        hidden_rows = frozenset(
            i + 1 for i, ri in sh.rowinfo_map.items() if ri.hidden
        )
        hidden_cols = frozenset(
            i + 1 for i, ci in sh.colinfo_map.items() if ci.hidden
        )
        rows, excel_rows = _finalize(rows)
        if not rows:
            continue
        grids.append(
            Grid(rows, excel_rows, merged, hidden_rows, hidden_cols, sh.name, filename)
        )
    return grids


def _xls_cell(cell, datemode) -> Cell:
    t = cell.ctype
    if t == xlrd.XL_CELL_TEXT:
        v = cell.value
    elif t == xlrd.XL_CELL_NUMBER:
        v = cell.value
    elif t == xlrd.XL_CELL_DATE:
        v = xldate_as_datetime(cell.value, datemode)
    elif t == xlrd.XL_CELL_BOOLEAN:
        v = bool(cell.value)
    elif t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        v = None
    else:  # XL_CELL_ERROR and any other -> keep raw
        v = cell.value
    return Cell(v, "General", False)


# --------------------------------------------------------------------------
# csv / txt (delimited text)
# --------------------------------------------------------------------------
def _load_csv(data: bytes, filename: str) -> list[Grid]:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        # UTF-16 with a byte-order mark (Excel "Unicode Text" export). The BOM
        # picks endianness; utf-8-sig would leave a NUL between every character
        # and the sniffer would see one garbage column.
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # default: comma
    # newline="" lets the csv module handle line endings itself (CRLF-safe,
    # and correct for newlines embedded in quoted fields).
    reader = csv.reader(StringIO(text, newline=""), dialect)
    rows = [
        [Cell(v if v != "" else None, "General", False) for v in row]
        for row in reader
    ]
    rows, excel_rows = _finalize(rows)
    return [Grid(rows, excel_rows, [], frozenset(), frozenset(), "(single)", filename)]


# --------------------------------------------------------------------------
# shared: trim trailing empties + rectangularize; return 1-based excel_rows
# --------------------------------------------------------------------------
def _finalize(rows: list) -> tuple:
    # Drop trailing all-empty rows (interior blanks are kept -- clean.py's job).
    while rows and all(c.value is None for c in rows[-1]):
        rows.pop()
    if not rows:
        return [], []
    # Rightmost column holding any value across all rows.
    last_col = 0
    for r in rows:
        for idx in range(len(r) - 1, -1, -1):
            if r[idx].value is not None:
                if idx + 1 > last_col:
                    last_col = idx + 1
                break
    # Make rectangular: pad short rows, trim trailing empty columns.
    for i, r in enumerate(rows):
        if len(r) < last_col:
            rows[i] = r + [Cell(None) for _ in range(last_col - len(r))]
        elif len(r) > last_col:
            rows[i] = r[:last_col]
    excel_rows = list(range(1, len(rows) + 1))
    return rows, excel_rows
